import asyncio
import io
import os
import logging
from collections import defaultdict
from datetime import datetime, timedelta, timezone

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import PieChart, BarChart, Reference

from aiogram import Bot, Dispatcher, F, BaseMiddleware
from aiogram.filters import CommandStart
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.fsm.storage.memory import MemoryStorage
from aiogram.fsm.storage.base import StorageKey
from aiogram.types import (
    Message, CallbackQuery,
    InlineKeyboardMarkup, InlineKeyboardButton,
    ReplyKeyboardMarkup, KeyboardButton,
    BufferedInputFile
)
from apscheduler.schedulers.asyncio import AsyncIOScheduler

from database import (
    init_db, init_settings_table, init_debts_tables,
    init_invest_tables, init_profit_table, init_export_history_table,
    init_accounts_table, init_view_log_table, init_access_requests_table,
    init_categories_table,
    migrate_existing_users,
    register_account, authenticate, get_account_by_id, get_account_by_tg,
    link_tg_to_account, unlink_tg, change_password, list_accounts,
    set_blocked, reset_password, set_role, get_account_role,
    log_view, get_views_of_user, get_views_by_admin,
    create_access_request, get_access_request, resolve_access_request,
    has_active_access, get_pending_requests_for_user, get_all_requests_for_user,
    add_operation, get_months, get_operations_for_month,
    get_all_operations_for_month, get_prev_month, get_month_totals,
    get_operations_for_range, get_last_operation_dt,
    get_earliest_operation_date,
    get_settings, save_settings,
    get_all_weekly_users, get_all_daily_users,
    get_all_reminder_users, update_reminder_last_sent,
    get_all_autoexport_users,
    add_debt, get_open_debts, get_all_debts, get_debt,
    get_all_debts_for_export, pay_debt, get_total_debt,
    update_debt, get_debt_payments, get_last_payment,
    rollback_last_payment, get_payments_all,
    add_investment, get_investments, get_investments_all,
    get_investments_for_month, get_investment_totals,
    get_investments_by_category, set_portfolio_value, get_portfolio_value,
    add_profit, get_total_realized_profit, get_profits,
    get_profits_all, get_profits_for_month, get_profits_by_category,
    add_export_history, get_export_history_filtered,
    get_export_history_by_period, get_export_by_id, delete_export,
    ensure_default_categories,
    get_categories, add_category, delete_category, rename_category,
)

# ============================================================
#                       НАСТРОЙКИ
# ============================================================
BOT_TOKEN = os.getenv("BOT_TOKEN")
EXPORTS_DIR = "exports"

WEEKDAYS_RU = ["Пн", "Вт", "Ср", "Чт", "Пт", "Сб", "Вс"]

CAT_KINDS = {
    "expense": "💸 Траты",
    "income": "💰 Зачисления",
    "invest": "📈 Инвестиции",
    "profit_source": "💵 Источники прибыли",
}

logging.basicConfig(level=logging.INFO)
bot = Bot(token=BOT_TOKEN)
_fsm_storage = MemoryStorage()
dp = Dispatcher(storage=_fsm_storage)
scheduler = AsyncIOScheduler()


# ============================================================
#                       FSM
# ============================================================
class AuthFlow(StatesGroup):
    login = State()
    password = State()
    register_login = State()
    register_password = State()


class ChangePwdFlow(StatesGroup):
    old = State()
    new = State()


class ExpenseFlow(StatesGroup):
    amount = State()
    category = State()
    comment = State()


class IncomeFlow(StatesGroup):
    amount = State()
    category = State()
    comment = State()
    debt_choice = State()
    debt_amount = State()


class DebtFlow(StatesGroup):
    amount = State()
    comment = State()


class DebtEditFlow(StatesGroup):
    amount = State()
    comment = State()


class WeeklyReportFlow(StatesGroup):
    time_input = State()
    tz_input = State()


class DailyReportFlow(StatesGroup):
    time_input = State()
    tz_input = State()


class ReminderFlow(StatesGroup):
    interval = State()
    silent_from = State()
    silent_to = State()


class AutoExportFlow(StatesGroup):
    weekday = State()
    time_input = State()
    period = State()


class InvestFlow(StatesGroup):
    amount = State()
    category = State()
    comment = State()


class InvestValueFlow(StatesGroup):
    value = State()


class ProfitFlow(StatesGroup):
    amount = State()
    category = State()
    source = State()
    comment = State()


class ExcelRangeFlow(StatesGroup):
    start_date = State()
    end_date = State()


class ExcelAutoFlow(StatesGroup):
    threshold = State()


class CatFlow(StatesGroup):
    choosing_kind = State()
    menu = State()
    add_name = State()
    rename_pick = State()
    rename_new = State()
    delete_pick = State()


# ============================================================
#                  АВТОРИЗАЦИЯ
# ============================================================
def current_account(tg_user_id: int):
    return get_account_by_tg(tg_user_id)


def uid(event) -> int:
    tg_id = event.from_user.id
    acc = current_account(tg_id)
    return acc[0] if acc else tg_id


class AuthMiddleware(BaseMiddleware):
    ALLOWED_CALLBACKS_PREFIX = ("auth:",)
    ALLOWED_CALLBACKS = set()

    AUTH_STATES = {
        AuthFlow.login.state,
        AuthFlow.password.state,
        AuthFlow.register_login.state,
        AuthFlow.register_password.state,
    }

    async def __call__(self, handler, event, data):
        tg_id = None
        if isinstance(event, Message):
            tg_id = event.from_user.id
            text = event.text or ""
            if text.startswith("/start"):
                return await handler(event, data)

            try:
                key = StorageKey(bot_id=bot.id, chat_id=event.chat.id,
                                 user_id=event.from_user.id)
                state = await _fsm_storage.get_state(key)
                if state in self.AUTH_STATES:
                    return await handler(event, data)
            except Exception:
                pass

        elif isinstance(event, CallbackQuery):
            tg_id = event.from_user.id
            if event.data and any(event.data.startswith(p)
                                  for p in self.ALLOWED_CALLBACKS_PREFIX):
                return await handler(event, data)
            if event.data in self.ALLOWED_CALLBACKS:
                return await handler(event, data)

        if tg_id is None:
            return await handler(event, data)

        acc = current_account(tg_id)
        if not acc:
            if isinstance(event, Message):
                await event.answer("🔒 Требуется вход. Нажмите /start.")
            else:
                await event.answer("🔒 Сначала войдите.", show_alert=True)
            return

        data["account_id"] = acc[0]
        data["account"] = acc
        return await handler(event, data)


dp.message.middleware(AuthMiddleware())
dp.callback_query.middleware(AuthMiddleware())


# ============================================================
#                    КЛАВИАТУРЫ
# ============================================================
def auth_menu_kb():
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🔑 Войти", callback_data="auth:login")],
        [InlineKeyboardButton(text="📝 Регистрация", callback_data="auth:register")],
    ])


def main_menu_kb(total_debt: float | None = None) -> ReplyKeyboardMarkup:
    if total_debt is None:
        debt_label = "💳 Задолженности"
    elif total_debt > 0:
        debt_label = f"💳 Задолженности • {total_debt:g} ₽"
    else:
        debt_label = "💳 Задолженности • нет долгов"

    return ReplyKeyboardMarkup(
        keyboard=[
            [KeyboardButton(text="💸 Записать трату")],
            [KeyboardButton(text="💰 Внести зачисление")],
            [KeyboardButton(text="📈 Инвестиции")],
            [KeyboardButton(text=debt_label)],
            [KeyboardButton(text="📊 История трат")],
            [KeyboardButton(text="📈 Сводка по категориям")],
            [KeyboardButton(text="🗓 Еженедельный отчёт")],
            [KeyboardButton(text="⏰ Ежедневный отчёт")],
            [KeyboardButton(text="🔔 Напоминания")],
            [KeyboardButton(text="📤 Экспорт в Excel")],
            [KeyboardButton(text="⚙️ Профиль")],
            [KeyboardButton(text="🚪 Выйти")],
        ],
        resize_keyboard=True
    )


def _menu_for(user_id: int) -> ReplyKeyboardMarkup:
    try:
        total = get_total_debt(user_id)
    except Exception:
        total = 0.0
    return main_menu_kb(total)


def category_kb(categories):
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text=c, callback_data=f"cat:{c}")]
        for c in categories
    ])


def skip_comment_kb():
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="Без комментария", callback_data="skip_comment")]
    ])


def months_kb(months, prefix: str = "hist"):
    buttons = []
    for m in months:
        y, mo = m.split("-")
        buttons.append([InlineKeyboardButton(text=f"{mo}.{y}",
                                             callback_data=f"{prefix}:{m}")])
    return InlineKeyboardMarkup(inline_keyboard=buttons)


def report_actions_kb(month: str):
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📋 Детализация",
                              callback_data=f"details:{month}")],
        [InlineKeyboardButton(text="📤 Экспорт в CSV",
                              callback_data=f"csv:{month}")],
    ])


def weekly_menu_kb(enabled, time_str, tz, send_excel=False):
    status = f"включён: вс {time_str} (UTC{tz:+d})" if enabled and time_str else "выключен"
    excel_status = "✅" if send_excel else "❌"
    buttons = [
        [InlineKeyboardButton(text=f"Статус: {status}", callback_data="noop")],
        [InlineKeyboardButton(text="🕐 Установить время", callback_data="weekly:set_time")],
        [InlineKeyboardButton(text="🌍 Часовой пояс", callback_data="weekly:set_tz")],
        [InlineKeyboardButton(text=f"📎 Excel в отчёте: {excel_status}",
                              callback_data="weekly:toggle_excel")],
    ]
    if enabled:
        buttons.append([InlineKeyboardButton(text="🔕 Выключить",
                                             callback_data="weekly:disable")])
    return InlineKeyboardMarkup(inline_keyboard=buttons)


def daily_menu_kb(enabled, time_str, tz, send_excel=False):
    status = f"включён: {time_str} (UTC{tz:+d})" if enabled and time_str else "выключен"
    excel_status = "✅" if send_excel else "❌"
    buttons = [
        [InlineKeyboardButton(text=f"Статус: {status}", callback_data="noop")],
        [InlineKeyboardButton(text="🕐 Установить время", callback_data="daily:set_time")],
        [InlineKeyboardButton(text="🌍 Часовой пояс", callback_data="daily:set_tz")],
        [InlineKeyboardButton(text=f"📎 Excel в отчёте: {excel_status}",
                              callback_data="daily:toggle_excel")],
    ]
    if enabled:
        buttons.append([InlineKeyboardButton(text="🔕 Выключить",
                                             callback_data="daily:disable")])
    return InlineKeyboardMarkup(inline_keyboard=buttons)


def reminder_menu_kb(s) -> InlineKeyboardMarkup:
    enabled = bool(s["reminder_enabled"])
    status = (f"включены: каждые {s['reminder_interval_hours']} ч"
              if enabled else "выключены")
    silent = f"{s['reminder_silent_from']:02d}:00–{s['reminder_silent_to']:02d}:00"
    buttons = [
        [InlineKeyboardButton(text=f"Статус: {status}", callback_data="noop")],
        [InlineKeyboardButton(text="⏱ Интервал", callback_data="reminder:interval")],
        [InlineKeyboardButton(text=f"🌙 Тихие часы: {silent}",
                              callback_data="reminder:silent_from")],
    ]
    if enabled:
        buttons.append([InlineKeyboardButton(text="🔕 Выключить",
                                             callback_data="reminder:disable")])
    else:
        buttons.append([InlineKeyboardButton(text="🔔 Включить",
                                             callback_data="reminder:enable")])
    return InlineKeyboardMarkup(inline_keyboard=buttons)


def debts_menu_kb():
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="➕ Добавить задолженность",
                              callback_data="debt:add")],
        [InlineKeyboardButton(text="📋 Список задолженностей",
                              callback_data="debt:list")],
    ])


def skip_debt_comment_kb():
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="Без пояснения", callback_data="debt:skip_comment")]
    ])


def debts_list_kb(debts, action_prefix: str = "debt:view"):
    buttons = []
    for debt_id, amount, remaining, comment, created_at in debts:
        cmt = (comment or "")[:25]
        label = f"#{debt_id} • {remaining:g} ₽ • {cmt}"
        buttons.append([InlineKeyboardButton(text=label,
                                             callback_data=f"{action_prefix}:{debt_id}")])
    buttons.append([InlineKeyboardButton(text="⬅️ Назад", callback_data="debt:menu")])
    return InlineKeyboardMarkup(inline_keyboard=buttons)


def debt_actions_kb(debt_id: int):
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📜 История платежей",
                              callback_data=f"debt:history:{debt_id}")],
        [InlineKeyboardButton(text="✏️ Изменить сумму",
                              callback_data=f"debt:edit_amount:{debt_id}")],
        [InlineKeyboardButton(text="✏️ Изменить комментарий",
                              callback_data=f"debt:edit_comment:{debt_id}")],
        [InlineKeyboardButton(text="↩️ Отменить последний платёж",
                              callback_data=f"debt:rollback:{debt_id}")],
        [InlineKeyboardButton(text="⬅️ К списку", callback_data="debt:list")],
    ])


def invest_menu_kb():
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📥 Пополнить", callback_data="invest:deposit")],
        [InlineKeyboardButton(text="📤 Вывести", callback_data="invest:withdraw")],
        [InlineKeyboardButton(text="💰 Зафиксировать прибыль",
                              callback_data="profit:add")],
        [InlineKeyboardButton(text="📊 Сводка", callback_data="invest:summary")],
        [InlineKeyboardButton(text="📋 Последние операции", callback_data="invest:list")],
        [InlineKeyboardButton(text="📜 История прибылей", callback_data="profit:list")],
        [InlineKeyboardButton(text="💼 Текущая стоимость портфеля",
                              callback_data="invest:set_value")],
    ])


def skip_invest_comment_kb():
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="Без комментария", callback_data="invest:skip_comment")]
    ])


def profit_source_kb(sources):
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text=s, callback_data=f"psrc:{s}")]
        for s in sources
    ])


def skip_profit_comment_kb():
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="Без комментария", callback_data="profit:skip_comment")]
    ])


def excel_menu_kb(s):
    def mark(v): return "✅" if v else "❌"
    period_map = {"week": "неделя", "month": "месяц", "all": "всё время"}
    if s["autoexport_enabled"] and s["autoexport_time"]:
        wd = WEEKDAYS_RU[s["autoexport_weekday"]]
        period = period_map.get(s["autoexport_period"], "?")
        auto_status = f"кажд. {wd} {s['autoexport_time']} ({period})"
    else:
        auto_status = "выключен"

    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📄 Выгрузить за месяц",
                              callback_data="xls:pick_month")],
        [InlineKeyboardButton(text="🗓 Экспорт за диапазон",
                              callback_data="xls:range")],
        [InlineKeyboardButton(text="─ Настройки листов ─", callback_data="noop")],
        [InlineKeyboardButton(text=f"{mark(s['excel_include_ops'])} Операции",
                              callback_data="xls:toggle:ops")],
        [InlineKeyboardButton(text=f"{mark(s['excel_include_debts'])} Задолженности",
                              callback_data="xls:toggle:debts")],
        [InlineKeyboardButton(text=f"{mark(s['excel_include_payments'])} Платежи по долгам",
                              callback_data="xls:toggle:payments")],
        [InlineKeyboardButton(text=f"{mark(s['excel_include_invest'])} Инвестиции",
                              callback_data="xls:toggle:invest")],
        [InlineKeyboardButton(text=f"{mark(s['excel_include_profits'])} Прибыли",
                              callback_data="xls:toggle:profits")],
        [InlineKeyboardButton(text="─ Авто-отправка при событиях ─",
                              callback_data="noop")],
        [InlineKeyboardButton(
            text=f"{mark(s['excel_on_debt_close'])} Excel при закрытии долга",
            callback_data="xls:auto:toggle:debt_close")],
        [InlineKeyboardButton(
            text=f"{mark(s['excel_on_big_op'])} Excel при крупной операции",
            callback_data="xls:auto:toggle:big_op")],
        [InlineKeyboardButton(
            text=f"⚙️ Порог крупной: {s['excel_big_op_threshold']:g} ₽",
            callback_data="xls:auto:big_op_threshold")],
        [InlineKeyboardButton(text="─ Автоэкспорт по расписанию ─",
                              callback_data="noop")],
        [InlineKeyboardButton(text=f"🗓 Автоэкспорт: {auto_status}",
                              callback_data="auto:xls:menu")],
        [InlineKeyboardButton(text="📜 История экспортов",
                              callback_data="xls:history")],
    ])


def autoexport_menu_kb(s):
    period_map = {"week": "неделя", "month": "месяц", "all": "всё время"}
    status = "выключен"
    if s["autoexport_enabled"] and s["autoexport_time"]:
        status = (f"включён: {WEEKDAYS_RU[s['autoexport_weekday']]} "
                  f"{s['autoexport_time']} ({period_map[s['autoexport_period']]})")
    buttons = [
        [InlineKeyboardButton(text=f"Статус: {status}", callback_data="noop")],
        [InlineKeyboardButton(text="📅 День недели", callback_data="auto:xls:weekday")],
        [InlineKeyboardButton(text="🕐 Время", callback_data="auto:xls:time")],
        [InlineKeyboardButton(text=f"📆 Период: {period_map[s['autoexport_period']]}",
                              callback_data="auto:xls:period")],
    ]
    if s["autoexport_enabled"]:
        buttons.append([InlineKeyboardButton(text="🔕 Выключить",
                                             callback_data="auto:xls:disable")])
    else:
        buttons.append([InlineKeyboardButton(text="🔔 Включить",
                                             callback_data="auto:xls:enable")])
    buttons.append([InlineKeyboardButton(text="⬅️ Назад", callback_data="xls:menu")])
    return InlineKeyboardMarkup(inline_keyboard=buttons)


REASON_LABELS = {
    "manual": "📄 Вручную",
    "weekly": "🗓 Недельный отчёт",
    "daily": "📅 Ежедневный отчёт",
    "autoexport": "🔄 Автоэкспорт",
    "debt_close": "💳 Закрытие долга",
    "big_op": "💸 Крупная операция",
    "range": "📆 Диапазон",
}


def history_filter_kb(current_filter):
    def mark(reason):
        return "• " if current_filter == reason else ""
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text=f"{mark(None)}Все",
                              callback_data="xls:hist:f:all")],
        [InlineKeyboardButton(text=f"{mark('manual')}📄 Вручную",
                              callback_data="xls:hist:f:manual"),
         InlineKeyboardButton(text=f"{mark('weekly')}🗓 Недельные",
                              callback_data="xls:hist:f:weekly")],
        [InlineKeyboardButton(text=f"{mark('daily')}📅 Ежедневные",
                              callback_data="xls:hist:f:daily"),
         InlineKeyboardButton(text=f"{mark('autoexport')}🔄 Авто",
                              callback_data="xls:hist:f:autoexport")],
        [InlineKeyboardButton(text=f"{mark('debt_close')}💳 Долги",
                              callback_data="xls:hist:f:debt_close"),
         InlineKeyboardButton(text=f"{mark('big_op')}💸 Крупные",
                              callback_data="xls:hist:f:big_op")],
        [InlineKeyboardButton(text=f"{mark('range')}📆 Диапазон",
                              callback_data="xls:hist:f:range")],
        [InlineKeyboardButton(text="📥 Экспорт истории в CSV",
                              callback_data="xls:hist:csv")],
        [InlineKeyboardButton(text="⬅️ Назад", callback_data="xls:menu")],
    ])


def admin_users_kb(role: str):
    buttons = [
        [InlineKeyboardButton(text="📋 Список аккаунтов", callback_data="admin:list")],
    ]
    buttons.append([InlineKeyboardButton(text="⬅️ Назад", callback_data="admin:back")])
    return InlineKeyboardMarkup(inline_keyboard=buttons)


def view_data_kb(target_id: int, role: str) -> InlineKeyboardMarkup:
    buttons = [
        [InlineKeyboardButton(text="📊 Сводка за месяц (агрегаты)",
                              callback_data=f"admin:view_month:{target_id}")],
        [InlineKeyboardButton(text="💳 Задолженности (агрегаты)",
                              callback_data=f"admin:view_debts:{target_id}")],
        [InlineKeyboardButton(text="📈 Инвестиции (агрегаты)",
                              callback_data=f"admin:view_inv:{target_id}")],
    ]
    if role == "admin":
        buttons.insert(1, [InlineKeyboardButton(
            text="📋 Все операции",
            callback_data=f"admin:view_ops:{target_id}")])
        buttons.append([InlineKeyboardButton(
            text="📜 История экспортов",
            callback_data=f"admin:view_exports:{target_id}")])
    buttons.append([InlineKeyboardButton(text="⬅️ К списку",
                                         callback_data="admin:list")])
    return InlineKeyboardMarkup(inline_keyboard=buttons)


def role_menu_kb(target_id: int):
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="👤 Сделать пользователем",
                              callback_data=f"admin:set_role:user:{target_id}")],
        [InlineKeyboardButton(text="🛡 Сделать модератором",
                              callback_data=f"admin:set_role:moderator:{target_id}")],
        [InlineKeyboardButton(text="👑 Сделать админом",
                              callback_data=f"admin:set_role:admin:{target_id}")],
        [InlineKeyboardButton(text="⬅️ Назад", callback_data="admin:list")],
    ])


def categories_main_kb():
    buttons = [
        [InlineKeyboardButton(text=label, callback_data=f"cat_menu:{kind}")]
        for kind, label in CAT_KINDS.items()
    ]
    buttons.append([InlineKeyboardButton(text="⬅️ Назад",
                                         callback_data="profile:back")])
    return InlineKeyboardMarkup(inline_keyboard=buttons)


def category_menu_kb(kind: str):
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="➕ Добавить",
                              callback_data=f"cat_add:{kind}")],
        [InlineKeyboardButton(text="✏️ Переименовать",
                              callback_data=f"cat_rename:{kind}")],
        [InlineKeyboardButton(text="🗑 Удалить",
                              callback_data=f"cat_delete:{kind}")],
        [InlineKeyboardButton(text="⬅️ Назад",
                              callback_data="cat_main")],
    ])
# ============================================================
#                        СТАРТ
# ============================================================
@dp.message(CommandStart())
async def cmd_start(message: Message, state: FSMContext):
    await state.clear()
    acc = current_account(message.from_user.id)
    if not acc:
        await message.answer(
            "👋 Привет!\n\nЭто бот для учёта финансов. "
            "Чтобы начать — войдите или зарегистрируйтесь.",
            reply_markup=auth_menu_kb()
        )
        return
    _, login, _, role, _, display_name, *_ = acc
    role_line = ""
    if role == "admin":
        role_line = "\n👑 Вы вошли как <b>администратор</b>."
    elif role == "moderator":
        role_line = "\n🛡 Вы вошли как <b>модератор</b>."
    await message.answer(
        f"👋 С возвращением, <b>{display_name or login}</b>!{role_line}\n\n"
        f"Выберите действие:",
        parse_mode="HTML",
        reply_markup=_menu_for(acc[0])
    )


# ============================================================
#                    ВХОД / РЕГИСТРАЦИЯ
# ============================================================
@dp.callback_query(F.data == "auth:login")
async def auth_login_start(call: CallbackQuery, state: FSMContext):
    await state.set_state(AuthFlow.login)
    await call.message.edit_text("Введите <b>логин</b>:", parse_mode="HTML")
    await call.answer()


@dp.message(AuthFlow.login)
async def auth_login_input(message: Message, state: FSMContext):
    await state.update_data(login=message.text.strip())
    await state.set_state(AuthFlow.password)
    await message.answer("Введите <b>пароль</b>:", parse_mode="HTML")


@dp.message(AuthFlow.password)
async def auth_password_input(message: Message, state: FSMContext):
    data = await state.get_data()
    login = data.get("login", "")
    password = message.text.strip()
    ok, msg, account_id = authenticate(login, password)
    await state.clear()
    if not ok:
        await message.answer(f"❌ {msg}\n\nПопробуйте снова:",
                             reply_markup=auth_menu_kb())
        return
    link_tg_to_account(account_id, message.from_user.id)
    acc = get_account_by_id(account_id)
    await message.answer(
        f"✅ Вход выполнен. Добро пожаловать, <b>{acc[5] or acc[1]}</b>!",
        parse_mode="HTML",
        reply_markup=_menu_for(account_id)
    )


@dp.callback_query(F.data == "auth:register")
async def auth_register_start(call: CallbackQuery, state: FSMContext):
    await state.set_state(AuthFlow.register_login)
    await call.message.edit_text(
        "📝 <b>Регистрация</b>\n\nПридумайте <b>логин</b> "
        "(мин. 3 символа):",
        parse_mode="HTML"
    )
    await call.answer()


@dp.message(AuthFlow.register_login)
async def auth_register_login(message: Message, state: FSMContext):
    login = message.text.strip()
    if len(login) < 3:
        await message.answer("❗ Логин не короче 3 символов.")
        return
    await state.update_data(reg_login=login)
    await state.set_state(AuthFlow.register_password)
    await message.answer("Придумайте <b>пароль</b> (мин. 6 символов):",
                         parse_mode="HTML")


@dp.message(AuthFlow.register_password)
async def auth_register_password(message: Message, state: FSMContext):
    password = message.text.strip()
    if len(password) < 6:
        await message.answer("❗ Пароль не короче 6 символов.")
        return
    data = await state.get_data()
    login = data.get("reg_login")
    ok, msg, account_id = register_account(login, password, display_name=login)
    await state.clear()
    if not ok:
        await message.answer(f"❌ {msg}", reply_markup=auth_menu_kb())
        return
    link_tg_to_account(account_id, message.from_user.id)
    await message.answer(
        f"✅ Аккаунт <b>{login}</b> создан.\nДобро пожаловать!",
        parse_mode="HTML",
        reply_markup=_menu_for(account_id)
    )


@dp.message(F.text == "🚪 Выйти")
async def logout(message: Message, state: FSMContext):
    await state.clear()
    unlink_tg(message.from_user.id)
    await message.answer(
        "🚪 Вы вышли из аккаунта.\n\nЧтобы войти снова — нажмите /start."
    )


# ============================================================
#                   ПРОФИЛЬ + СМЕНА ПАРОЛЯ
# ============================================================
def _profile_text_buttons(acc):
    acc_id, login, _, role, _, display_name, created_at, last_login_at = acc
    role_map = {"admin": "👑 Администратор",
                "moderator": "🛡 Модератор",
                "user": "👤 Пользователь"}
    pending = get_pending_requests_for_user(acc_id)
    lines = [
        f"⚙️ <b>Профиль</b>\n",
        f"👤 Логин: <b>{login}</b>",
        f"🏷 Имя: {display_name or '—'}",
        f"🎚 Роль: {role_map.get(role, role)}",
        f"📅 Создан: {created_at[:10]}",
        f"🕐 Последний вход: {(last_login_at or '—')[:16]}",
    ]
    if pending:
        lines.append(f"\n🔐 <b>Активных запросов: {len(pending)}</b>")

    buttons = [
        [InlineKeyboardButton(text="🔑 Сменить пароль",
                              callback_data="profile:change_pwd")],
        [InlineKeyboardButton(text="🏷 Мои категории",
                              callback_data="cat_main")],
        [InlineKeyboardButton(text="👀 Кто смотрел мои данные",
                              callback_data="profile:views_of_me")],
        [InlineKeyboardButton(
            text="🔐 Запросы на просмотр" +
                 (f" ({len(pending)})" if pending else ""),
            callback_data="profile:requests")],
    ]
    if role in ("admin", "moderator"):
        buttons.append([InlineKeyboardButton(
            text="👥 Управление пользователями",
            callback_data="admin:users")])
        buttons.append([InlineKeyboardButton(
            text="📖 Мои просмотры",
            callback_data="profile:my_views")])
    return "\n".join(lines), InlineKeyboardMarkup(inline_keyboard=buttons)


@dp.message(F.text == "⚙️ Профиль")
async def profile(message: Message, state: FSMContext):
    await state.clear()
    acc = current_account(message.from_user.id)
    if not acc:
        await message.answer("Сессия истекла. Нажмите /start.")
        return
    text, kb = _profile_text_buttons(acc)
    await message.answer(text, parse_mode="HTML", reply_markup=kb)


@dp.callback_query(F.data == "profile:back")
async def profile_back(call: CallbackQuery):
    acc = current_account(call.from_user.id)
    if not acc:
        await call.answer("Сессия истекла.", show_alert=True)
        return
    text, kb = _profile_text_buttons(acc)
    await call.message.edit_text(text, parse_mode="HTML", reply_markup=kb)
    await call.answer()


@dp.callback_query(F.data == "profile:change_pwd")
async def change_pwd_start(call: CallbackQuery, state: FSMContext):
    await state.set_state(ChangePwdFlow.old)
    await call.message.edit_text("Введите <b>текущий пароль</b>:",
                                 parse_mode="HTML")
    await call.answer()


@dp.message(ChangePwdFlow.old)
async def change_pwd_old(message: Message, state: FSMContext):
    await state.update_data(old=message.text.strip())
    await state.set_state(ChangePwdFlow.new)
    await message.answer("Введите <b>новый пароль</b> (мин. 6):",
                         parse_mode="HTML")


@dp.message(ChangePwdFlow.new)
async def change_pwd_new(message: Message, state: FSMContext):
    acc = current_account(message.from_user.id)
    if not acc:
        await message.answer("Сессия истекла.")
        await state.clear()
        return
    data = await state.get_data()
    old_pwd = data.get("old", "")
    new_pwd = message.text.strip()
    ok, msg = change_password(acc[0], old_pwd, new_pwd)
    await state.clear()
    if not ok:
        await message.answer(f"❌ {msg}")
        return
    await message.answer(f"✅ {msg}", reply_markup=_menu_for(acc[0]))


# ============================================================
#              МЕНЮ ПОЛЬЗОВАТЕЛЬСКИХ КАТЕГОРИЙ
# ============================================================
@dp.callback_query(F.data == "cat_main")
async def categories_main(call: CallbackQuery, state: FSMContext):
    await state.clear()
    acc = current_account(call.from_user.id)
    if not acc:
        await call.answer("Сессия истекла.", show_alert=True)
        return
    ensure_default_categories(acc[0])
    await call.message.edit_text(
        "🏷 <b>Мои категории</b>\n\nВыберите раздел:",
        parse_mode="HTML",
        reply_markup=categories_main_kb())
    await call.answer()


@dp.callback_query(F.data.startswith("cat_menu:"))
async def category_open(call: CallbackQuery, state: FSMContext):
    kind = call.data.split(":", 1)[1]
    if kind not in CAT_KINDS:
        await call.answer("Неизвестный раздел.", show_alert=True)
        return
    acc = current_account(call.from_user.id)
    if not acc:
        await call.answer("Сессия истекла.", show_alert=True)
        return
    names = get_categories(acc[0], kind)
    lines = [f"🏷 <b>{CAT_KINDS[kind]}</b>\n"]
    if not names:
        lines.append("<i>Пока нет ни одной категории.</i>")
    else:
        for i, name in enumerate(names, 1):
            lines.append(f"{i}. {name}")
    await call.message.edit_text("\n".join(lines), parse_mode="HTML",
                                 reply_markup=category_menu_kb(kind))
    await call.answer()


@dp.callback_query(F.data.startswith("cat_add:"))
async def category_add_start(call: CallbackQuery, state: FSMContext):
    kind = call.data.split(":", 1)[1]
    await state.set_state(CatFlow.add_name)
    await state.update_data(cat_kind=kind)
    await call.message.edit_text(
        f"Введите название новой категории для «{CAT_KINDS.get(kind, kind)}»:"
    )
    await call.answer()


@dp.message(CatFlow.add_name)
async def category_add_input(message: Message, state: FSMContext):
    acc = current_account(message.from_user.id)
    if not acc:
        await message.answer("Сессия истекла.")
        await state.clear()
        return
    data = await state.get_data()
    kind = data.get("cat_kind")
    ok, msg = add_category(acc[0], kind, message.text)
    await state.clear()
    if not ok:
        await message.answer(f"❌ {msg}")
        return
    names = get_categories(acc[0], kind)
    lines = [f"✅ {msg}\n", f"🏷 <b>{CAT_KINDS[kind]}</b>\n"]
    for i, name in enumerate(names, 1):
        lines.append(f"{i}. {name}")
    await message.answer("\n".join(lines), parse_mode="HTML",
                         reply_markup=category_menu_kb(kind))


@dp.callback_query(F.data.startswith("cat_rename:"))
async def category_rename_start(call: CallbackQuery, state: FSMContext):
    kind = call.data.split(":", 1)[1]
    acc = current_account(call.from_user.id)
    if not acc:
        await call.answer("Сессия истекла.", show_alert=True)
        return
    names = get_categories(acc[0], kind)
    if not names:
        await call.answer("Нет категорий.", show_alert=True)
        return
    await state.set_state(CatFlow.rename_pick)
    await state.update_data(cat_kind=kind)
    buttons = [[InlineKeyboardButton(text=n, callback_data=f"cat_rp:{n}")]
               for n in names]
    buttons.append([InlineKeyboardButton(text="⬅️ Назад",
                                         callback_data=f"cat_menu:{kind}")])
    await call.message.edit_text(
        "Какую категорию переименовать?",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons))
    await call.answer()


@dp.callback_query(CatFlow.rename_pick, F.data.startswith("cat_rp:"))
async def category_rename_pick(call: CallbackQuery, state: FSMContext):
    old_name = call.data.split(":", 1)[1]
    await state.update_data(old_name=old_name)
    await state.set_state(CatFlow.rename_new)
    await call.message.edit_text(f"Новое название для «{old_name}»:")
    await call.answer()


@dp.message(CatFlow.rename_new)
async def category_rename_input(message: Message, state: FSMContext):
    acc = current_account(message.from_user.id)
    if not acc:
        await message.answer("Сессия истекла.")
        await state.clear()
        return
    data = await state.get_data()
    kind = data.get("cat_kind")
    old_name = data.get("old_name")
    ok, msg = rename_category(acc[0], kind, old_name, message.text)
    await state.clear()
    if not ok:
        await message.answer(f"❌ {msg}")
        return
    names = get_categories(acc[0], kind)
    lines = [f"✅ {msg}\n", f"🏷 <b>{CAT_KINDS[kind]}</b>\n"]
    for i, name in enumerate(names, 1):
        lines.append(f"{i}. {name}")
    await message.answer("\n".join(lines), parse_mode="HTML",
                         reply_markup=category_menu_kb(kind))


@dp.callback_query(F.data.startswith("cat_delete:"))
async def category_delete_start(call: CallbackQuery, state: FSMContext):
    kind = call.data.split(":", 1)[1]
    acc = current_account(call.from_user.id)
    if not acc:
        await call.answer("Сессия истекла.", show_alert=True)
        return
    names = get_categories(acc[0], kind)
    if len(names) <= 1:
        await call.answer("Нельзя удалить последнюю категорию.", show_alert=True)
        return
    await state.set_state(CatFlow.delete_pick)
    await state.update_data(cat_kind=kind)
    buttons = [[InlineKeyboardButton(text=n, callback_data=f"cat_dp:{n}")]
               for n in names]
    buttons.append([InlineKeyboardButton(text="⬅️ Назад",
                                         callback_data=f"cat_menu:{kind}")])
    await call.message.edit_text(
        "Какую категорию удалить?",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons))
    await call.answer()


@dp.callback_query(CatFlow.delete_pick, F.data.startswith("cat_dp:"))
async def category_delete_confirm(call: CallbackQuery, state: FSMContext):
    name = call.data.split(":", 1)[1]
    acc = current_account(call.from_user.id)
    if not acc:
        await call.answer("Сессия истекла.", show_alert=True)
        return
    data = await state.get_data()
    kind = data.get("cat_kind")
    ok, msg = delete_category(acc[0], kind, name)
    await state.clear()
    if not ok:
        await call.answer(f"❌ {msg}", show_alert=True)
        return
    names = get_categories(acc[0], kind)
    lines = [f"✅ {msg}\n", f"🏷 <b>{CAT_KINDS[kind]}</b>\n"]
    for i, n in enumerate(names, 1):
        lines.append(f"{i}. {n}")
    await call.message.edit_text("\n".join(lines), parse_mode="HTML",
                                 reply_markup=category_menu_kb(kind))
    await call.answer("Удалено")


# ============================================================
#              КТО СМОТРЕЛ МОИ ДАННЫЕ / МОИ ПРОСМОТРЫ
# ============================================================
@dp.callback_query(F.data == "profile:views_of_me")
async def views_of_me(call: CallbackQuery):
    acc = current_account(call.from_user.id)
    if not acc:
        await call.answer("Сессия истекла.", show_alert=True)
        return
    rows = get_views_of_user(acc[0], limit=30)
    if not rows:
        await call.message.edit_text(
            "👀 <b>Кто смотрел мои данные</b>\n\n<i>Пока никто.</i>",
            parse_mode="HTML",
            reply_markup=InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(text="⬅️ К профилю",
                                      callback_data="profile:back")]]))
        await call.answer()
        return
    lines = ["👀 <b>Кто смотрел мои данные</b>\n"]
    for vid, viewer_id, viewer_login, section, detail, created_at in rows:
        login_disp = viewer_login or f"#{viewer_id}"
        tail = f" ({detail})" if detail else ""
        lines.append(f"• {created_at[:16]} | <b>{login_disp}</b> | {section}{tail}")
    text = "\n".join(lines)
    if len(text) > 4000:
        text = text[:4000] + "\n…"
    await call.message.edit_text(
        text, parse_mode="HTML",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="⬅️ К профилю",
                                  callback_data="profile:back")]]))
    await call.answer()


@dp.callback_query(F.data == "profile:my_views")
async def my_views(call: CallbackQuery):
    acc = current_account(call.from_user.id)
    if not acc or acc[3] not in ("admin", "moderator"):
        await call.answer("Только для админа/модератора.", show_alert=True)
        return
    rows = get_views_by_admin(acc[0], limit=30)
    if not rows:
        await call.message.edit_text(
            "📖 <b>Мои просмотры</b>\n\n<i>Вы никого не смотрели.</i>",
            parse_mode="HTML",
            reply_markup=InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(text="⬅️ К профилю",
                                      callback_data="profile:back")]]))
        await call.answer()
        return
    lines = ["📖 <b>Мои просмотры</b>\n"]
    for vid, target_id, target_login, section, detail, created_at in rows:
        login_disp = target_login or f"#{target_id}"
        tail = f" ({detail})" if detail else ""
        lines.append(f"• {created_at[:16]} | <b>{login_disp}</b> | {section}{tail}")
    text = "\n".join(lines)
    if len(text) > 4000:
        text = text[:4000] + "\n…"
    await call.message.edit_text(
        text, parse_mode="HTML",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="⬅️ К профилю",
                                  callback_data="profile:back")]]))
    await call.answer()


# ============================================================
#              ЗАПРОСЫ НА ПРОСМОТР — СПИСОК
# ============================================================
@dp.callback_query(F.data == "profile:requests")
async def profile_requests(call: CallbackQuery):
    acc = current_account(call.from_user.id)
    if not acc:
        await call.answer("Сессия истекла.", show_alert=True)
        return
    rows = get_all_requests_for_user(acc[0], limit=30)
    if not rows:
        await call.message.edit_text(
            "🔐 <b>Запросы на просмотр</b>\n\n<i>Запросов нет.</i>",
            parse_mode="HTML",
            reply_markup=InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(text="⬅️ К профилю",
                                      callback_data="profile:back")]]))
        await call.answer()
        return
    status_map = {"pending": "⏳ Ожидает", "approved": "✅ Одобрен",
                  "denied": "❌ Отклонён", "expired": "⌛ Истёк"}
    lines = ["🔐 <b>Запросы на просмотр</b>\n"]
    for req_id, viewer_id, viewer_login, section, status, created, resolved, expires in rows:
        st = status_map.get(status, status)
        login_disp = viewer_login or f"#{viewer_id}"
        lines.append(f"• {created[:16]} | <b>{login_disp}</b> | {section} | {st}")
        if status == "approved" and expires:
            lines.append(f"  <i>до {expires[:16]}</i>")
    text = "\n".join(lines)
    if len(text) > 4000:
        text = text[:4000] + "\n…"

    pending_rows = [r for r in rows if r[4] == "pending"]
    action_rows = []
    for req_id, viewer_id, viewer_login, section, *_ in pending_rows[:5]:
        action_rows.append([
            InlineKeyboardButton(text=f"✅ #{req_id}",
                                 callback_data=f"access:approve:{req_id}"),
            InlineKeyboardButton(text=f"❌ #{req_id}",
                                 callback_data=f"access:deny:{req_id}"),
        ])
    action_rows.append([InlineKeyboardButton(text="⬅️ К профилю",
                                             callback_data="profile:back")])
    await call.message.edit_text(
        text, parse_mode="HTML",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=action_rows))
    await call.answer()


# ============================================================
#              ЗАПРОСЫ — APPROVE / DENY / RECHECK
# ============================================================
@dp.callback_query(F.data.startswith("access:approve:"))
async def access_approve(call: CallbackQuery):
    req_id = int(call.data.split(":")[2])
    req = get_access_request(req_id)
    if not req or req[4] != "pending":
        await call.answer("Запрос уже неактуален.", show_alert=True)
        return
    acc = current_account(call.from_user.id)
    if not acc or acc[0] != req[2]:
        await call.answer("Это не ваш запрос.", show_alert=True)
        return
    resolve_access_request(req_id, approve=True, minutes=30)
    viewer = get_account_by_id(req[1])
    if viewer and viewer[2]:
        try:
            await bot.send_message(
                viewer[2],
                f"✅ <b>Доступ разрешён</b>\n\n"
                f"Пользователь <b>{acc[1]}</b> разрешил просмотр раздела "
                f"<b>{req[3]}</b> на 30 минут.",
                parse_mode="HTML")
        except Exception as e:
            logging.warning(f"Не удалось уведомить модератора: {e}")
    await call.message.edit_text(
        f"✅ Вы разрешили просмотр раздела <b>{req[3]}</b> на 30 минут.",
        parse_mode="HTML")
    await call.answer("Разрешено")


@dp.callback_query(F.data.startswith("access:deny:"))
async def access_deny(call: CallbackQuery):
    req_id = int(call.data.split(":")[2])
    req = get_access_request(req_id)
    if not req or req[4] != "pending":
        await call.answer("Запрос уже неактуален.", show_alert=True)
        return
    acc = current_account(call.from_user.id)
    if not acc or acc[0] != req[2]:
        await call.answer("Это не ваш запрос.", show_alert=True)
        return
    resolve_access_request(req_id, approve=False)
    viewer = get_account_by_id(req[1])
    if viewer and viewer[2]:
        try:
            await bot.send_message(
                viewer[2],
                f"❌ <b>Доступ отклонён</b>\n\n"
                f"Пользователь <b>{acc[1]}</b> отклонил запрос на просмотр "
                f"раздела <b>{req[3]}</b>.",
                parse_mode="HTML")
        except Exception as e:
            logging.warning(f"Не удалось уведомить модератора: {e}")
    await call.message.edit_text(
        f"❌ Вы отклонили запрос на просмотр раздела <b>{req[3]}</b>.",
        parse_mode="HTML")
    await call.answer("Отклонено")


@dp.callback_query(F.data.startswith("access:recheck:"))
async def access_recheck(call: CallbackQuery):
    parts = call.data.split(":", 3)
    target_id = int(parts[2])
    acc = current_account(call.from_user.id)
    if not acc or acc[3] not in ("admin", "moderator"):
        await call.answer("Нет доступа.", show_alert=True)
        return
    if has_active_access(acc[0], target_id):
        await call.answer("✅ Доступ одобрен! Откройте раздел заново.",
                          show_alert=True)
        return
    await call.answer("⏳ Пока нет разрешения.", show_alert=True)


# ============================================================
#              ЛОГ ПРОСМОТРА + УВЕДОМЛЕНИЕ
# ============================================================
async def _log_view_and_notify(viewer_acc, target_id: int,
                               section: str, detail: str = None):
    log_view(viewer_acc[0], target_id, section, detail)
    if target_id == viewer_acc[0]:
        return
    target = get_account_by_id(target_id)
    if not target:
        return
    target_tg = target[2]
    if not target_tg:
        return
    role_ru = "администратором" if viewer_acc[3] == "admin" else "модератором"
    try:
        await bot.send_message(
            target_tg,
            f"🔔 <b>Уведомление</b>\n\n"
            f"Ваша статистика просмотрена {role_ru}.\n"
            f"Раздел: <b>{section}</b>"
            + (f"\nДетали: {detail}" if detail else ""),
            parse_mode="HTML")
    except Exception as e:
        logging.warning(f"Не удалось уведомить {target_tg}: {e}")


async def _try_access_or_ask(viewer_acc, target_id: int,
                             section: str, call: CallbackQuery) -> bool:
    viewer_id = viewer_acc[0]
    role = viewer_acc[3]
    if viewer_id == target_id:
        return True
    if role == "admin":
        await _log_view_and_notify(viewer_acc, target_id, section)
        return True
    if has_active_access(viewer_id, target_id):
        await _log_view_and_notify(viewer_acc, target_id, section)
        return True
    req_id = create_access_request(viewer_id, target_id, section)
    target = get_account_by_id(target_id)
    if target and target[2]:
        try:
            await bot.send_message(
                target[2],
                f"🔐 <b>Запрос на просмотр</b>\n\n"
                f"<b>{viewer_acc[1]}</b> (модератор) запрашивает доступ "
                f"к разделу <b>{section}</b>.\n\nРазрешить на 30 минут?",
                parse_mode="HTML",
                reply_markup=InlineKeyboardMarkup(inline_keyboard=[
                    [InlineKeyboardButton(text="✅ Разрешить",
                                          callback_data=f"access:approve:{req_id}")],
                    [InlineKeyboardButton(text="❌ Отклонить",
                                          callback_data=f"access:deny:{req_id}")]]))
        except Exception as e:
            logging.warning(f"Не удалось отправить запрос {target[2]}: {e}")
    await call.message.edit_text(
        f"🔐 <b>Запрос отправлен</b>\n\n"
        f"Пользователь <b>{target[1] if target else target_id}</b> получил "
        f"запрос на просмотр раздела <b>{section}</b>.\n\n"
        f"Разрешение действует 30 минут после одобрения.",
        parse_mode="HTML",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="🔄 Обновить",
                                  callback_data=f"access:recheck:{target_id}:{section}")],
            [InlineKeyboardButton(text="⬅️ К списку",
                                  callback_data="admin:list")]]))
    return False


# ============================================================
#                   АДМИН-ПАНЕЛЬ
# ============================================================
@dp.callback_query(F.data == "admin:users")
async def admin_users(call: CallbackQuery, state: FSMContext):
    acc = current_account(call.from_user.id)
    if not acc or acc[3] not in ("admin", "moderator"):
        await call.answer("Доступ только для админа/модератора.", show_alert=True)
        return
    role = acc[3]
    role_emoji = "👑" if role == "admin" else "🛡"
    await call.message.edit_text(
        f"{role_emoji} <b>Управление пользователями</b>\n\n"
        f"Ваша роль: <b>{role}</b>",
        parse_mode="HTML",
        reply_markup=admin_users_kb(role))
    await call.answer()


@dp.callback_query(F.data == "admin:back")
async def admin_back(call: CallbackQuery, state: FSMContext):
    await state.clear()
    acc = current_account(call.from_user.id)
    if not acc:
        await call.answer("Сессия истекла.", show_alert=True)
        return
    text, kb = _profile_text_buttons(acc)
    await call.message.edit_text(text, parse_mode="HTML", reply_markup=kb)
    await call.answer()


@dp.callback_query(F.data == "admin:list")
async def admin_list(call: CallbackQuery):
    acc = current_account(call.from_user.id)
    if not acc or acc[3] not in ("admin", "moderator"):
        await call.answer("Нет доступа.", show_alert=True)
        return
    role = acc[3]
    rows = list_accounts()
    role_map = {"admin": "👑", "moderator": "🛡", "user": "👤"}
    lines = ["👥 <b>Все аккаунты</b>\n"]
    for acc_id, login, tg_id, r, is_blocked, name, created, last_login in rows:
        emoji = role_map.get(r, "👤")
        block = " 🚫" if is_blocked else ""
        lines.append(f"<b>#{acc_id}</b> {login} {emoji}{block}\n"
                     f"  TG: {tg_id or '—'} | создан: {created[:10]}")
    text = "\n".join(lines)
    if len(text) > 4000:
        text = text[:4000] + "\n…"

    action_rows = []
    for acc_id, login, *_ in rows[:10]:
        if acc_id == acc[0]:
            continue
        line = [InlineKeyboardButton(text=f"#{acc_id} 👁 Смотреть",
                                     callback_data=f"admin:view:{acc_id}")]
        if role == "admin":
            line.append(InlineKeyboardButton(
                text=f"#{acc_id} 🎚 Роль",
                callback_data=f"admin:role_menu:{acc_id}"))
        action_rows.append(line)

    action_rows.append([InlineKeyboardButton(text="⬅️ Назад",
                                             callback_data="admin:users")])
    await call.message.edit_text(text, parse_mode="HTML",
                                 reply_markup=InlineKeyboardMarkup(
                                     inline_keyboard=action_rows))
    await call.answer()


@dp.callback_query(F.data.startswith("admin:view:"))
async def admin_view(call: CallbackQuery, state: FSMContext):
    acc = current_account(call.from_user.id)
    if not acc or acc[3] not in ("admin", "moderator"):
        await call.answer("Нет доступа.", show_alert=True)
        return
    target_id = int(call.data.split(":")[2])
    target = get_account_by_id(target_id)
    if not target:
        await call.answer("Аккаунт не найден.", show_alert=True)
        return
    allowed = await _try_access_or_ask(acc, target_id, "Профиль пользователя", call)
    if not allowed:
        await call.answer()
        return
    _, login, tg_id, r, is_blocked, name, created, last_login = target
    role_map = {"admin": "👑 Админ", "moderator": "🛡 Модератор", "user": "👤 Пользователь"}
    text = (
        f"👁 <b>Просмотр данных аккаунта #{target_id}</b>\n\n"
        f"👤 Логин: <b>{login}</b>\n"
        f"🎚 Роль: {role_map.get(r, r)}\n"
        f"📅 Создан: {created[:10]}\n"
        f"🕐 Последний вход: {(last_login or '—')[:16]}\n")
    if acc[3] == "moderator":
        text += "\n🛡 <i>Ограниченный доступ: только агрегаты.</i>\n"
    await call.message.edit_text(text, parse_mode="HTML",
                                 reply_markup=view_data_kb(target_id, acc[3]))
    await call.answer()


@dp.callback_query(F.data.startswith("admin:role_menu:"))
async def admin_role_menu(call: CallbackQuery):
    acc = current_account(call.from_user.id)
    if not acc or acc[3] != "admin":
        await call.answer("Только админ.", show_alert=True)
        return
    target_id = int(call.data.split(":")[2])
    target = get_account_by_id(target_id)
    if not target:
        await call.answer("Не найден.", show_alert=True)
        return
    await call.message.edit_text(
        f"🎚 <b>Роль для аккаунта #{target_id} ({target[1]})</b>\n\n"
        f"Текущая: <b>{target[3]}</b>",
        parse_mode="HTML",
        reply_markup=role_menu_kb(target_id))
    await call.answer()


@dp.callback_query(F.data.startswith("admin:set_role:"))
async def admin_set_role(call: CallbackQuery):
    acc = current_account(call.from_user.id)
    if not acc or acc[3] != "admin":
        await call.answer("Только админ.", show_alert=True)
        return
    parts = call.data.split(":")
    new_role = parts[2]
    target_id = int(parts[3])
    if target_id == acc[0] and new_role != "admin":
        await call.answer("Нельзя снять с себя права админа.", show_alert=True)
        return
    set_role(target_id, new_role)
    await call.answer(f"Роль изменена на {new_role}")
    await admin_role_menu(call)
# ============================================================
#                        ТРАТЫ
# ============================================================
@dp.message(F.text == "💸 Записать трату")
async def start_expense(message: Message, state: FSMContext):
    await state.clear()
    await state.set_state(ExpenseFlow.amount)
    await message.answer("Введите сумму траты (например, 1500 или 1500.50):")


@dp.message(ExpenseFlow.amount)
async def expense_amount(message: Message, state: FSMContext):
    try:
        amount = float(message.text.replace(",", ".").strip())
        if amount <= 0:
            raise ValueError
    except ValueError:
        await message.answer("❗ Введите корректное положительное число.")
        return
    await state.update_data(amount=amount)
    await state.set_state(ExpenseFlow.category)
    cats = get_categories(uid(message), "expense")
    await message.answer("Выберите категорию траты:",
                         reply_markup=category_kb(cats))


@dp.callback_query(ExpenseFlow.category, F.data.startswith("cat:"))
async def expense_category(call: CallbackQuery, state: FSMContext):
    category = call.data.split(":", 1)[1]
    await state.update_data(category=category)
    await state.set_state(ExpenseFlow.comment)
    await call.message.edit_text(
        f"Категория: {category}\n\n"
        f"Напишите комментарий или нажмите «Без комментария»:",
        reply_markup=skip_comment_kb())
    await call.answer()


@dp.message(ExpenseFlow.comment)
async def expense_comment(message: Message, state: FSMContext):
    await _save_operation(message, state, op_type="expense")


@dp.callback_query(ExpenseFlow.comment, F.data == "skip_comment")
async def expense_skip_comment(call: CallbackQuery, state: FSMContext):
    await _save_operation_callback(call, state, op_type="expense")


# ============================================================
#                    ЗАЧИСЛЕНИЯ
# ============================================================
@dp.message(F.text == "💰 Внести зачисление")
async def start_income(message: Message, state: FSMContext):
    await state.clear()
    await state.set_state(IncomeFlow.amount)
    await message.answer("Введите сумму зачисления:")


@dp.message(IncomeFlow.amount)
async def income_amount(message: Message, state: FSMContext):
    try:
        amount = float(message.text.replace(",", ".").strip())
        if amount <= 0:
            raise ValueError
    except ValueError:
        await message.answer("❗ Введите корректное положительное число.")
        return
    await state.update_data(amount=amount)
    await state.set_state(IncomeFlow.category)
    cats = get_categories(uid(message), "income")
    await message.answer("Выберите категорию зачисления:",
                         reply_markup=category_kb(cats))


@dp.callback_query(IncomeFlow.category, F.data.startswith("cat:"))
async def income_category(call: CallbackQuery, state: FSMContext):
    category = call.data.split(":", 1)[1]
    await state.update_data(category=category)
    debts = get_open_debts(uid(call))
    if debts:
        await state.set_state(IncomeFlow.debt_choice)
        await call.message.edit_text(
            f"Категория: {category}\n\nПогасить одну из задолженностей?",
            reply_markup=InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(text="💳 Да, погасить долг",
                                      callback_data="income:paydebt")],
                [InlineKeyboardButton(text="➡️ Нет, продолжить",
                                      callback_data="income:skip_debt")]]))
    else:
        await state.set_state(IncomeFlow.comment)
        await call.message.edit_text(
            f"Категория: {category}\n\n"
            f"Напишите комментарий или нажмите «Без комментария»:",
            reply_markup=skip_comment_kb())
    await call.answer()


@dp.callback_query(IncomeFlow.debt_choice, F.data == "income:paydebt")
async def income_paydebt_menu(call: CallbackQuery, state: FSMContext):
    debts = get_open_debts(uid(call))
    await state.set_state(IncomeFlow.debt_amount)
    await call.message.edit_text("Выберите задолженность для погашения:",
                                 reply_markup=debts_list_kb(debts))
    await call.answer()


@dp.callback_query(IncomeFlow.debt_choice, F.data == "income:skip_debt")
async def income_skip_debt(call: CallbackQuery, state: FSMContext):
    data = await state.get_data()
    await state.set_state(IncomeFlow.comment)
    await call.message.edit_text(
        f"Категория: {data['category']}\n\n"
        f"Напишите комментарий или нажмите «Без комментария»:",
        reply_markup=skip_comment_kb())
    await call.answer()


@dp.callback_query(IncomeFlow.debt_amount, F.data.startswith("debt:view:"))
async def income_debt_selected(call: CallbackQuery, state: FSMContext):
    debt_id = int(call.data.split(":")[2])
    debt = get_debt(debt_id, uid(call))
    if not debt:
        await call.answer("Долг не найден.", show_alert=True)
        return
    await state.update_data(debt_id=debt_id)
    await call.message.edit_text(
        f"Долг #{debt_id}\nОстаток: <b>{debt[2]:g}</b> ₽\n\n"
        f"Введите сумму погашения (не больше остатка):",
        parse_mode="HTML")
    await call.answer()


@dp.message(IncomeFlow.debt_amount)
async def income_debt_amount(message: Message, state: FSMContext):
    data = await state.get_data()
    debt_id = data.get("debt_id")
    if not debt_id:
        await message.answer("Что-то пошло не так.")
        await state.clear()
        return
    debt = get_debt(debt_id, uid(message))
    if not debt:
        await message.answer("Долг не найден.")
        await state.clear()
        return
    remaining = debt[2]
    try:
        amount = float(message.text.replace(",", ".").strip())
        if amount <= 0:
            raise ValueError
    except ValueError:
        await message.answer("❗ Введите корректное положительное число.")
        return
    if amount > remaining:
        await message.answer(f"❗ Сумма больше остатка ({remaining:g} ₽).")
        return
    income_amount = data["amount"]
    income_category = data["category"]
    add_operation(user_id=uid(message), op_type="income",
                  amount=income_amount, category=income_category,
                  comment=f"Погашение задолженности #{debt_id}")
    new_remaining = pay_debt(debt_id, uid(message), amount)
    await state.clear()
    if new_remaining == 0:
        status = "✅ Долг полностью погашен!"
        s = get_settings(uid(message))
        if s["excel_on_debt_close"]:
            try:
                today = datetime.now().strftime("%Y-%m-%d")
                await _send_excel(uid(message), today, today,
                                  title=f"📎 Закрытие долга #{debt_id} — {today}",
                                  reason="debt_close",
                                  filename=f"debt_closed_{debt_id}_{today}.xlsx")
            except Exception as e:
                logging.warning(f"Excel при закрытии долга: {e}")
    else:
        status = f"Остаток по долгу #{debt_id}: <b>{new_remaining:g}</b> ₽"
    await message.answer(
        f"✅ Зачисление <b>{income_amount:g}</b> ₽ записано.\n"
        f"💳 Погашено по долгу #{debt_id}: <b>{amount:g}</b> ₽\n{status}",
        parse_mode="HTML",
        reply_markup=_menu_for(uid(message)))


@dp.message(IncomeFlow.comment)
async def income_comment(message: Message, state: FSMContext):
    await _save_operation(message, state, op_type="income")


@dp.callback_query(IncomeFlow.comment, F.data == "skip_comment")
async def income_skip_comment(call: CallbackQuery, state: FSMContext):
    await _save_operation_callback(call, state, op_type="income")


# ============================================================
#                СОХРАНЕНИЕ ОПЕРАЦИЙ
# ============================================================
async def _save_operation(message: Message, state: FSMContext, op_type: str):
    data = await state.get_data()
    comment = message.text.strip()
    await _finalize_save(uid(message), data, comment, op_type)
    await message.answer("✅ Записано!", reply_markup=_menu_for(uid(message)))
    await state.clear()


async def _save_operation_callback(call: CallbackQuery, state: FSMContext, op_type: str):
    data = await state.get_data()
    await _finalize_save(uid(call), data, "—", op_type)
    await call.message.edit_text("✅ Записано!")
    await call.message.answer("Главное меню:", reply_markup=_menu_for(uid(call)))
    await call.answer()
    await state.clear()


async def _finalize_save(user_id: int, data: dict, comment: str, op_type: str):
    add_operation(user_id=user_id, op_type=op_type, amount=data["amount"],
                  category=data["category"], comment=comment)
    s = get_settings(user_id)
    if s["excel_on_big_op"] and data["amount"] >= s["excel_big_op_threshold"]:
        try:
            today = datetime.now().strftime("%Y-%m-%d")
            await _send_excel(user_id, today, today,
                              title=f"📎 Крупная операция {data['amount']:g} ₽ — {today}",
                              reason="big_op",
                              filename=f"big_op_{today}.xlsx")
        except Exception as e:
            logging.warning(f"Excel при крупной операции: {e}")


# ============================================================
#               АГРЕГАЦИЯ / ОТЧЁТЫ / ГРАФИК
# ============================================================
def _aggregate(ops):
    expense_sums = defaultdict(float)
    expense_counts = defaultdict(int)
    income_sums = defaultdict(float)
    income_counts = defaultdict(int)
    total_expense = 0.0
    total_income = 0.0
    expense_lines = []
    income_lines = []
    for op_type, amount, category, comment, created_at in ops:
        dt = datetime.strptime(created_at, "%Y-%m-%d %H:%M:%S").strftime("%d.%m %H:%M")
        if op_type == "expense":
            total_expense += amount
            expense_sums[category] += amount
            expense_counts[category] += 1
            expense_lines.append(f"  • {dt} | <b>{amount:g}</b> ₽ | {category} | {comment}")
        else:
            total_income += amount
            income_sums[category] += amount
            income_counts[category] += 1
            income_lines.append(f"  • {dt} | <b>{amount:g}</b> ₽ | {category} | {comment}")
    return {
        "expense_sums": expense_sums, "expense_counts": expense_counts,
        "income_sums": income_sums, "income_counts": income_counts,
        "total_expense": total_expense, "total_income": total_income,
        "expense_lines": expense_lines, "income_lines": income_lines,
    }


def _category_sort_key(cat, categories_order):
    try:
        return (0, categories_order.index(cat))
    except ValueError:
        return (1, cat)


def _diff_str(current: float, prev: float) -> str:
    if prev == 0:
        if current == 0:
            return "0 ₽ (без изменений)"
        return f"{current:g} ₽ (в прошлом периоде не было)"
    diff = current - prev
    pct = (diff / prev) * 100
    sign = "+" if diff >= 0 else ""
    if diff > 0:
        arrow = "🔺"
    elif diff < 0:
        arrow = "🔻"
    else:
        arrow = "➖"
    return f"{arrow} {sign}{pct:.1f}% (было {prev:g} ₽, стало {current:g} ₽)"


def _top_expenses(ops, limit=3):
    rows = []
    for op_type, amount, category, comment, created_at in ops:
        if op_type == "expense":
            rows.append((amount, category, comment, created_at))
    rows.sort(key=lambda r: r[0], reverse=True)
    return rows[:limit]


def _make_report(month: str, agg, prev_totals=None,
                 expense_order=None, income_order=None) -> str:
    expense_order = expense_order or []
    income_order = income_order or []
    y, mo = month.split("-")
    lines = [f"📅 <b>Отчёт за {mo}.{y}</b>\n"]
    if agg["expense_sums"]:
        lines.append("💸 <b>Траты по категориям:</b>")
        for cat in sorted(agg["expense_sums"].keys(),
                          key=lambda c: _category_sort_key(c, expense_order)):
            s = agg["expense_sums"][cat]
            cnt = agg["expense_counts"][cat]
            percent = (s / agg["total_expense"] * 100) if agg["total_expense"] else 0
            lines.append(f"  • <b>{cat}</b>: {s:g} ₽  ({cnt} шт., {percent:.1f}%)")
        lines.append(f"  <i>Итого трат: {agg['total_expense']:g} ₽</i>\n")
    if agg["income_sums"]:
        lines.append("💰 <b>Зачисления по категориям:</b>")
        for cat in sorted(agg["income_sums"].keys(),
                          key=lambda c: _category_sort_key(c, income_order)):
            s = agg["income_sums"][cat]
            cnt = agg["income_counts"][cat]
            percent = (s / agg["total_income"] * 100) if agg["total_income"] else 0
            lines.append(f"  • <b>{cat}</b>: {s:g} ₽  ({cnt} шт., {percent:.1f}%)")
        lines.append(f"  <i>Итого зачислений: {agg['total_income']:g} ₽</i>\n")
    if prev_totals is not None:
        prev_exp, prev_inc = prev_totals
        if prev_exp > 0 or prev_inc > 0:
            lines.append("📈 <b>Сравнение с прошлым месяцем:</b>")
            lines.append(f"  • Траты: {_diff_str(agg['total_expense'], prev_exp)}")
            lines.append(f"  • Зачисления: {_diff_str(agg['total_income'], prev_inc)}\n")
    lines.append("━━━━━━━━━━━━━━━")
    lines.append(f"💸 Итого трат: <b>{agg['total_expense']:g}</b> ₽")
    lines.append(f"💰 Итого зачислений: <b>{agg['total_income']:g}</b> ₽")
    lines.append(f"📊 Баланс: <b>{agg['total_income'] - agg['total_expense']:g}</b> ₽")
    return "\n".join(lines)


def _make_weekly_report(start_date, end_date, agg, ops=None,
                        prev_totals=None, expense_order=None,
                        income_order=None):
    expense_order = expense_order or []
    income_order = income_order or []
    lines = [f"🗓 <b>Отчёт за неделю</b>\n"
             f"<i>{start_date.strftime('%d.%m')} — {end_date.strftime('%d.%m.%Y')}</i>\n"]
    if not agg["expense_sums"] and not agg["income_sums"]:
        lines.append("За эту неделю не было ни трат, ни зачислений.")
        return "\n".join(lines)
    if agg["expense_sums"]:
        lines.append("💸 <b>Траты по категориям:</b>")
        for cat in sorted(agg["expense_sums"].keys(),
                          key=lambda c: _category_sort_key(c, expense_order)):
            s = agg["expense_sums"][cat]
            cnt = agg["expense_counts"][cat]
            percent = (s / agg["total_expense"] * 100) if agg["total_expense"] else 0
            lines.append(f"  • <b>{cat}</b>: {s:g} ₽  ({cnt} шт., {percent:.1f}%)")
        lines.append(f"  <i>Итого: {agg['total_expense']:g} ₽</i>\n")
    if agg["income_sums"]:
        lines.append("💰 <b>Зачисления по категориям:</b>")
        for cat in sorted(agg["income_sums"].keys(),
                          key=lambda c: _category_sort_key(c, income_order)):
            s = agg["income_sums"][cat]
            cnt = agg["income_counts"][cat]
            lines.append(f"  • <b>{cat}</b>: {s:g} ₽ ({cnt} шт.)")
        lines.append(f"  <i>Итого: {agg['total_income']:g} ₽</i>\n")
    if prev_totals is not None:
        prev_exp, prev_inc = prev_totals
        if prev_exp > 0 or prev_inc > 0:
            lines.append("📈 <b>Сравнение с прошлой неделей:</b>")
            lines.append(f"  • Траты: {_diff_str(agg['total_expense'], prev_exp)}")
            lines.append(f"  • Зачисления: {_diff_str(agg['total_income'], prev_inc)}\n")
    top_expenses = _top_expenses(ops, limit=3) if ops else []
    if top_expenses:
        lines.append("🔥 <b>Крупнейшие траты недели:</b>")
        for i, (amount, category, comment, created_at) in enumerate(top_expenses, 1):
            dt = datetime.strptime(created_at, "%Y-%m-%d %H:%M:%S").strftime("%d.%m")
            cmt = comment if comment and comment != "—" else ""
            tail = f" — {cmt}" if cmt else ""
            lines.append(f"  {i}. <b>{amount:g}</b> ₽ | {category} | {dt}{tail}")
        lines.append("")
    lines.append("━━━━━━━━━━━━━━━")
    lines.append(f"💸 Итого трат: <b>{agg['total_expense']:g}</b> ₽")
    lines.append(f"💰 Итого зачислений: <b>{agg['total_income']:g}</b> ₽")
    lines.append(f"📊 Баланс: <b>{agg['total_income'] - agg['total_expense']:g}</b> ₽")
    return "\n".join(lines)


def _make_daily_report(date_str: str, agg, prev_totals=None,
                       expense_order=None, income_order=None) -> str:
    expense_order = expense_order or []
    income_order = income_order or []
    dt = datetime.strptime(date_str, "%Y-%m-%d")
    lines = [f"📅 <b>Отчёт за {dt.strftime('%d.%m.%Y')}</b>\n"]
    if not agg["expense_sums"] and not agg["income_sums"]:
        lines.append("За этот день не было ни трат, ни зачислений.")
        return "\n".join(lines)
    if agg["expense_sums"]:
        lines.append("💸 <b>Траты:</b>")
        for cat in sorted(agg["expense_sums"].keys(),
                          key=lambda c: _category_sort_key(c, expense_order)):
            lines.append(f"  • {cat}: <b>{agg['expense_sums'][cat]:g}</b> ₽")
        lines.append(f"  <i>Итого: {agg['total_expense']:g} ₽</i>\n")
    if agg["income_sums"]:
        lines.append("💰 <b>Зачисления:</b>")
        for cat in sorted(agg["income_sums"].keys(),
                          key=lambda c: _category_sort_key(c, income_order)):
            lines.append(f"  • {cat}: <b>{agg['income_sums'][cat]:g}</b> ₽")
        lines.append(f"  <i>Итого: {agg['total_income']:g} ₽</i>\n")
    lines.append("━━━━━━━━━━━━━━━")
    lines.append(f"💸 Итого трат: <b>{agg['total_expense']:g}</b> ₽")
    lines.append(f"💰 Итого зачислений: <b>{agg['total_income']:g}</b> ₽")
    lines.append(f"📊 Баланс: <b>{agg['total_income'] - agg['total_expense']:g}</b> ₽")
    return "\n".join(lines)


def _build_weekly_chart(start_date, end_date, ops):
    days = [(start_date + timedelta(days=i)) for i in range(7)]
    day_labels = [WEEKDAYS_RU[d.weekday()] for d in days]
    expense_by_day = [0.0] * 7
    income_by_day = [0.0] * 7
    for op_type, amount, category, comment, created_at in ops:
        try:
            dt = datetime.strptime(created_at, "%Y-%m-%d %H:%M:%S")
        except ValueError:
            continue
        idx = (dt.date() - start_date.date()).days
        if 0 <= idx < 7:
            if op_type == "expense":
                expense_by_day[idx] += amount
            else:
                income_by_day[idx] += amount
    if all(v == 0 for v in expense_by_day) and all(v == 0 for v in income_by_day):
        return None
    cumulative = []
    running = 0.0
    for e, i in zip(expense_by_day, income_by_day):
        running += (i - e)
        cumulative.append(running)
    fig, ax1 = plt.subplots(figsize=(8, 4), dpi=120)
    x = range(7)
    bw = 0.4
    ax1.bar([i - bw / 2 for i in x], expense_by_day, width=bw,
            color="#e74c3c", label="Траты")
    ax1.bar([i + bw / 2 for i in x], income_by_day, width=bw,
            color="#2ecc71", label="Зачисления")
    ax1.set_xticks(list(x))
    ax1.set_xticklabels(day_labels)
    ax1.set_ylabel("Сумма, ₽")
    ax1.grid(axis="y", linestyle="--", alpha=0.4)
    ax1.legend(loc="upper left", fontsize=9)
    ax2 = ax1.twinx()
    ax2.plot(list(x), cumulative, color="#3498db", marker="o",
             linewidth=2, label="Баланс")
    ax2.set_ylabel("Баланс, ₽", color="#3498db")
    ax2.tick_params(axis="y", labelcolor="#3498db")
    ax2.axhline(0, color="#3498db", linewidth=0.5, alpha=0.4)
    ax2.legend(loc="upper right", fontsize=9)
    ax1.set_title(f"Неделя {start_date.strftime('%d.%m')} — {end_date.strftime('%d.%m.%Y')}",
                  fontsize=11)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return buf.read()


# ============================================================
#              EXCEL — ОБЩИЙ КОНСТРУКТОР
# ============================================================
def _build_period_excel(user_id: int, start_date: str, end_date: str,
                        title: str) -> bytes:
    wb = Workbook()
    s = get_settings(user_id)
    include_ops = bool(s["excel_include_ops"])
    include_debts = bool(s["excel_include_debts"])
    include_payments = bool(s["excel_include_payments"])
    include_invest = bool(s["excel_include_invest"])
    include_profits = bool(s["excel_include_profits"])

    expense_order = get_categories(user_id, "expense")
    income_order = get_categories(user_id, "income")

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="34495E")
    title_font = Font(bold=True, size=13, color="2C3E50")
    bold = Font(bold=True)
    thin = Side(border_style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")

    def style_header(ws, row: int, col_count: int):
        for c in range(1, col_count + 1):
            cell = ws.cell(row=row, column=c)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = border

    def style_body(ws, start_row: int, end_row: int, col_count: int):
        for r in range(start_row, end_row + 1):
            for c in range(1, col_count + 1):
                cell = ws.cell(row=r, column=c)
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    def autosize(ws):
        for col in ws.columns:
            max_len = 0
            letter = get_column_letter(col[0].column)
            for cell in col:
                v = cell.value
                if v is None:
                    continue
                length = max(len(str(x)) for x in str(v).split("\n"))
                if length > max_len:
                    max_len = length
            ws.column_dimensions[letter].width = min(max_len + 3, 50)

    ops = get_operations_for_range(user_id, start_date, end_date)
    invs = [i for i in get_investments_all(user_id)
            if start_date <= i[5][:10] <= end_date]
    profits = [p for p in get_profits_all(user_id)
               if start_date <= p[5][:10] <= end_date]
    payments = [p for p in get_payments_all(user_id)
                if start_date <= p[3][:10] <= end_date]
    all_debts = get_all_debts_for_export(user_id)

    total_exp = sum(o[1] for o in ops if o[0] == "expense")
    total_inc = sum(o[1] for o in ops if o[0] == "income")
    total_inv_in = sum(i[2] for i in invs if i[1] == "deposit")
    total_inv_out = sum(i[2] for i in invs if i[1] == "withdraw")
    total_profit = sum(p[1] for p in profits)
    total_paid_debts = sum(p[2] for p in payments)

    exp_by_cat = defaultdict(float)
    for o in ops:
        if o[0] == "expense":
            exp_by_cat[o[2]] += o[1]
    inc_by_cat = defaultdict(float)
    for o in ops:
        if o[0] == "income":
            inc_by_cat[o[2]] += o[1]

    ws = wb.active
    ws.title = "Сводка"
    ws["A1"] = title
    ws["A1"].font = title_font

    row = 3
    ws.cell(row=row, column=1, value="Показатель").font = bold
    ws.cell(row=row, column=2, value="Значение").font = bold
    style_header(ws, row, 2)
    row += 1
    summary_rows = [
        ("💸 Траты за период", round(total_exp, 2)),
        ("💰 Зачисления за период", round(total_inc, 2)),
        ("📊 Баланс (доходы − траты)", round(total_inc - total_exp, 2)),
        ("", ""),
        ("📥 Инвестиции: вложено", round(total_inv_in, 2)),
        ("📤 Инвестиции: выведено", round(total_inv_out, 2)),
        ("💰 Зафиксированная прибыль", round(total_profit, 2)),
        ("", ""),
        ("💳 Погашено долгов", round(total_paid_debts, 2)),
        ("💳 Активных долгов (на сейчас)",
         round(sum(d[2] for d in all_debts if d[4] == 0), 2)),
    ]
    for label, value in summary_rows:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=value)
        row += 1

    row += 2
    ws.cell(row=row, column=1, value="💸 Траты по категориям").font = bold
    row += 1
    ws.cell(row=row, column=1, value="Категория").font = bold
    ws.cell(row=row, column=2, value="Сумма, ₽").font = bold
    style_header(ws, row, 2)
    row += 1
    exp_data_start = row
    for cat, val in sorted(exp_by_cat.items(),
                           key=lambda x: _category_sort_key(x[0], expense_order)):
        ws.cell(row=row, column=1, value=cat)
        ws.cell(row=row, column=2, value=round(val, 2))
        row += 1
    exp_data_end = row - 1

    row += 2
    ws.cell(row=row, column=1, value="💰 Зачисления по категориям").font = bold
    row += 1
    ws.cell(row=row, column=1, value="Категория").font = bold
    ws.cell(row=row, column=2, value="Сумма, ₽").font = bold
    style_header(ws, row, 2)
    row += 1
    inc_data_start = row
    for cat, val in sorted(inc_by_cat.items(),
                           key=lambda x: _category_sort_key(x[0], income_order)):
        ws.cell(row=row, column=1, value=cat)
        ws.cell(row=row, column=2, value=round(val, 2))
        row += 1
    inc_data_end = row - 1

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 16

    if exp_data_end >= exp_data_start:
        pie = PieChart()
        pie.title = "Траты по категориям"
        labels = Reference(ws, min_col=1, min_row=exp_data_start, max_row=exp_data_end)
        data = Reference(ws, min_col=2, min_row=exp_data_start - 1, max_row=exp_data_end)
        pie.add_data(data, titles_from_data=True)
        pie.set_categories(labels)
        pie.height = 8
        pie.width = 12
        ws.add_chart(pie, "D3")

    if inc_data_end >= inc_data_start:
        bar = BarChart()
        bar.type = "col"
        bar.title = "Зачисления по категориям"
        bar.y_axis.title = "₽"
        bar.x_axis.title = "Категория"
        labels = Reference(ws, min_col=1, min_row=inc_data_start, max_row=inc_data_end)
        data = Reference(ws, min_col=2, min_row=inc_data_start - 1, max_row=inc_data_end)
        bar.add_data(data, titles_from_data=True)
        bar.set_categories(labels)
        bar.height = 8
        bar.width = 12
        ws.add_chart(bar, "D20")

    if include_ops:
        ws2 = wb.create_sheet("Операции")
        headers = ["Дата", "Тип", "Сумма, ₽", "Категория", "Комментарий"]
        ws2.append(headers)
        style_header(ws2, 1, len(headers))
        for op_type, amount, category, comment, created_at in ops:
            ws2.append([
                created_at,
                "Трата" if op_type == "expense" else "Зачисление",
                round(amount, 2), category, comment or ""])
        if ops:
            last = ws2.max_row + 1
            ws2.cell(row=last, column=2, value="ИТОГО ТРАТ").font = bold
            ws2.cell(row=last, column=3, value=round(total_exp, 2)).font = bold
            last += 1
            ws2.cell(row=last, column=2, value="ИТОГО ЗАЧИСЛЕНИЙ").font = bold
            ws2.cell(row=last, column=3, value=round(total_inc, 2)).font = bold
        style_body(ws2, 2, ws2.max_row, len(headers))
        autosize(ws2)
        ws2.freeze_panes = "A2"

    if include_debts:
        ws3 = wb.create_sheet("Задолженности")
        headers = ["ID", "Дата создания", "Сумма, ₽", "Остаток, ₽",
                   "Комментарий", "Статус", "Дата закрытия"]
        ws3.append(headers)
        style_header(ws3, 1, len(headers))
        for d_id, amount, remaining, comment, is_closed, created_at, closed_at in all_debts:
            ws3.append([
                d_id, created_at, round(amount, 2), round(remaining, 2),
                comment or "", "Закрыт" if is_closed else "Активен",
                closed_at or ""])
        style_body(ws3, 2, ws3.max_row, len(headers))
        autosize(ws3)
        ws3.freeze_panes = "A2"

    if include_payments:
        ws4 = wb.create_sheet("Платежи по долгам")
        headers = ["ID долга", "Дата", "Сумма, ₽", "Комментарий долга"]
        ws4.append(headers)
        style_header(ws4, 1, len(headers))
        for p_id, debt_id, amount, created_at, debt_comment in payments:
            ws4.append([debt_id, created_at, round(amount, 2), debt_comment or ""])
        style_body(ws4, 2, ws4.max_row, len(headers))
        autosize(ws4)
        ws4.freeze_panes = "A2"

    if include_invest:
        ws5 = wb.create_sheet("Инвестиции")
        headers = ["Дата", "Тип", "Сумма, ₽", "Категория", "Комментарий"]
        ws5.append(headers)
        style_header(ws5, 1, len(headers))
        for i_id, inv_type, amount, category, comment, created_at in invs:
            ws5.append([
                created_at,
                "Пополнение" if inv_type == "deposit" else "Вывод",
                round(amount, 2), category, comment or ""])
        if invs:
            last = ws5.max_row + 1
            ws5.cell(row=last, column=2, value="ИТОГО ВЛОЖЕНО").font = bold
            ws5.cell(row=last, column=3, value=round(total_inv_in, 2)).font = bold
            last += 1
            ws5.cell(row=last, column=2, value="ИТОГО ВЫВЕДЕНО").font = bold
            ws5.cell(row=last, column=3, value=round(total_inv_out, 2)).font = bold
        style_body(ws5, 2, ws5.max_row, len(headers))
        autosize(ws5)
        ws5.freeze_panes = "A2"

    if include_profits:
        ws6 = wb.create_sheet("Прибыли")
        headers = ["Дата", "Сумма, ₽", "Категория", "Источник", "Комментарий"]
        ws6.append(headers)
        style_header(ws6, 1, len(headers))
        for p_id, amount, category, source, comment, created_at in profits:
            ws6.append([
                created_at, round(amount, 2), category,
                source or "", comment or ""])
        if profits:
            last = ws6.max_row + 1
            ws6.cell(row=last, column=1, value="ИТОГО ПРИБЫЛЬ").font = bold
            ws6.cell(row=last, column=2, value=round(total_profit, 2)).font = bold
        style_body(ws6, 2, ws6.max_row, len(headers))
        autosize(ws6)
        ws6.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


async def _send_excel(user_id: int, start_date: str, end_date: str,
                      title: str, reason: str,
                      filename: str | None = None) -> int:
    data = _build_period_excel(user_id, start_date, end_date, title=title)
    if filename is None:
        filename = f"finance_{start_date}_to_{end_date}.xlsx"
    user_dir = os.path.join(EXPORTS_DIR, str(user_id))
    os.makedirs(user_dir, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    disk_name = f"{ts}_{filename}"
    filepath = os.path.join(user_dir, disk_name)
    with open(filepath, "wb") as f:
        f.write(data)
    export_id = add_export_history(
        user_id=user_id, filename=filename, reason=reason,
        period_start=start_date, period_end=end_date,
        size_bytes=len(data), filepath=filepath)
    try:
        file = BufferedInputFile(data, filename=filename)
        await bot.send_document(user_id, document=file, caption=title)
    except Exception as e:
        logging.warning(f"Не удалось отправить Excel {user_id}: {e}")
        raise
    return export_id
# ============================================================
#                    ИСТОРИЯ ТРАТ
# ============================================================
@dp.message(F.text == "📊 История трат")
async def history(message: Message, state: FSMContext):
    await state.clear()
    months = get_months(uid(message))
    if not months:
        await message.answer("Пока нет ни одной записи.")
        return
    await message.answer("Выберите месяц:", reply_markup=months_kb(months, "hist"))


@dp.message(F.text == "📈 Сводка по категориям")
async def summary_menu(message: Message, state: FSMContext):
    await state.clear()
    months = get_months(uid(message))
    if not months:
        await message.answer("Пока нет ни одной записи.")
        return
    await message.answer("Выберите месяц для сводки:",
                         reply_markup=months_kb(months, "summary"))


@dp.callback_query(F.data.startswith("hist:"))
async def history_month(call: CallbackQuery):
    month = call.data.split(":", 1)[1]
    user_id = uid(call)
    ops = get_operations_for_month(user_id, month)
    if not ops:
        await call.message.edit_text("За этот месяц записей нет.")
        await call.answer()
        return
    agg = _aggregate(ops)
    prev_month = get_prev_month(month)
    prev_exp, prev_inc = get_month_totals(user_id, prev_month)
    text = _make_report(
        month, agg,
        prev_totals=(prev_exp, prev_inc),
        expense_order=get_categories(user_id, "expense"),
        income_order=get_categories(user_id, "income"),
    )
    if len(text) > 4000:
        text = text[:4000] + "\n… (обрезано)"
    await call.message.edit_text(text, parse_mode="HTML",
                                 reply_markup=report_actions_kb(month))
    await call.answer()


@dp.callback_query(F.data.startswith("summary:"))
async def summary_month(call: CallbackQuery):
    month = call.data.split(":", 1)[1]
    user_id = uid(call)
    ops = get_operations_for_month(user_id, month)
    agg = _aggregate(ops)
    exp_order = get_categories(user_id, "expense")
    inc_order = get_categories(user_id, "income")
    y, mo = month.split("-")
    lines = [f"📈 <b>Сводка за {mo}.{y}</b>\n"]
    if agg["expense_sums"]:
        lines.append("💸 <b>Траты:</b>")
        for cat in sorted(agg["expense_sums"].keys(),
                          key=lambda c: _category_sort_key(c, exp_order)):
            lines.append(f"  • {cat}: <b>{agg['expense_sums'][cat]:g}</b> ₽")
        lines.append("")
    if agg["income_sums"]:
        lines.append("💰 <b>Зачисления:</b>")
        for cat in sorted(agg["income_sums"].keys(),
                          key=lambda c: _category_sort_key(c, inc_order)):
            lines.append(f"  • {cat}: <b>{agg['income_sums'][cat]:g}</b> ₽")
        lines.append("")
    lines.append(f"📊 Баланс: <b>{agg['total_income'] - agg['total_expense']:g}</b> ₽")
    await call.message.edit_text("\n".join(lines), parse_mode="HTML")
    await call.answer()


@dp.callback_query(F.data.startswith("details:"))
async def show_details(call: CallbackQuery):
    month = call.data.split(":", 1)[1]
    ops = get_operations_for_month(uid(call), month)
    agg = _aggregate(ops)
    y, mo = month.split("-")
    lines = [f"📋 <b>Детализация за {mo}.{y}</b>\n"]
    if agg["expense_lines"]:
        lines.append("💸 <b>Траты:</b>")
        lines.extend(agg["expense_lines"])
        lines.append("")
    if agg["income_lines"]:
        lines.append("💰 <b>Зачисления:</b>")
        lines.extend(agg["income_lines"])
        lines.append("")
    text = "\n".join(lines) or "Нет записей."
    if len(text) > 4000:
        text = text[:4000] + "\n… (обрезано)"
    await call.message.answer(text, parse_mode="HTML")
    await call.answer()


@dp.callback_query(F.data.startswith("csv:"))
async def export_csv(call: CallbackQuery):
    month = call.data.split(":", 1)[1]
    ops = get_operations_for_month(uid(call), month)
    if not ops:
        await call.answer("Нет данных.", show_alert=True)
        return
    lines = ["Дата;Тип;Сумма;Категория;Комментарий"]
    for op_type, amount, category, comment, created_at in ops:
        type_ru = "Трата" if op_type == "expense" else "Зачисление"
        safe_comment = (comment or "").replace('"', '""')
        if ";" in safe_comment or '"' in safe_comment or "\n" in safe_comment:
            safe_comment = f'"{safe_comment}"'
        lines.append(f"{created_at};{type_ru};{amount:g};{category};{safe_comment}")
    csv_data = "\n".join(lines).encode("utf-8-sig")
    file = BufferedInputFile(csv_data, filename=f"report_{month}.csv")
    await call.message.answer_document(file, caption=f"📤 Экспорт за {month}")
    await call.answer()


# ============================================================
#                       ЗАДОЛЖЕННОСТИ
# ============================================================
@dp.message(F.text.startswith("💳 Задолженности"))
async def debts_menu(message: Message, state: FSMContext):
    await state.clear()
    total = get_total_debt(uid(message))
    await message.answer(
        f"💳 <b>Задолженности</b>\n\n"
        f"Общая сумма активных долгов: <b>{total:g}</b> ₽",
        parse_mode="HTML", reply_markup=debts_menu_kb())


@dp.callback_query(F.data == "debt:menu")
async def debts_menu_cb(call: CallbackQuery, state: FSMContext):
    await state.clear()
    total = get_total_debt(uid(call))
    await call.message.edit_text(
        f"💳 <b>Задолженности</b>\n\n"
        f"Общая сумма активных долгов: <b>{total:g}</b> ₽",
        parse_mode="HTML", reply_markup=debts_menu_kb())
    await call.answer()


@dp.callback_query(F.data == "debt:add")
async def debt_add_start(call: CallbackQuery, state: FSMContext):
    await state.set_state(DebtFlow.amount)
    await call.message.edit_text("Введите сумму задолженности:")
    await call.answer()


@dp.message(DebtFlow.amount)
async def debt_add_amount(message: Message, state: FSMContext):
    try:
        amount = float(message.text.replace(",", ".").strip())
        if amount <= 0:
            raise ValueError
    except ValueError:
        await message.answer("❗ Введите корректное положительное число.")
        return
    await state.update_data(amount=amount)
    await state.set_state(DebtFlow.comment)
    await message.answer(
        "Напишите пояснение (откуда задолженность) "
        "или нажмите «Без пояснения»:",
        reply_markup=skip_debt_comment_kb())


@dp.message(DebtFlow.comment)
async def debt_add_comment(message: Message, state: FSMContext):
    data = await state.get_data()
    debt_id = add_debt(uid(message), data["amount"], message.text.strip())
    await state.clear()
    await message.answer(
        f"✅ Задолженность #{debt_id} добавлена: <b>{data['amount']:g}</b> ₽",
        parse_mode="HTML", reply_markup=_menu_for(uid(message)))


@dp.callback_query(DebtFlow.comment, F.data == "debt:skip_comment")
async def debt_add_skip(call: CallbackQuery, state: FSMContext):
    data = await state.get_data()
    debt_id = add_debt(uid(call), data["amount"], "—")
    await state.clear()
    await call.message.edit_text(
        f"✅ Задолженность #{debt_id} добавлена: <b>{data['amount']:g}</b> ₽",
        parse_mode="HTML")
    await call.message.answer("Главное меню:", reply_markup=_menu_for(uid(call)))
    await call.answer()


@dp.callback_query(F.data == "debt:list")
async def debt_list(call: CallbackQuery, state: FSMContext):
    await state.clear()
    all_debts = get_all_debts(uid(call))
    if not all_debts:
        await call.message.edit_text("Пока нет ни одной задолженности.",
                                     reply_markup=debts_menu_kb())
        await call.answer()
        return
    total_active = sum(d[2] for d in all_debts if d[4] == 0)
    lines = [f"📋 <b>Список задолженностей:</b>",
             f"<i>Общий остаток по активным: {total_active:g} ₽</i>\n"]
    for debt_id, amount, remaining, comment, is_closed, created_at, closed_at in all_debts:
        status = "✅ закрыт" if is_closed else f"остаток: <b>{remaining:g}</b> ₽"
        cmt = comment if comment and comment != "—" else "—"
        lines.append(f"<b>#{debt_id}</b> — {amount:g} ₽ | {status}\n"
                     f"  💬 {cmt}\n  📅 {created_at[:10]}")
    text = "\n".join(lines)
    if len(text) > 4000:
        text = text[:4000] + "\n… (обрезано)"
    active = [d for d in all_debts if d[4] == 0]
    kb = debts_list_kb(active) if active else debts_menu_kb()
    await call.message.edit_text(text, parse_mode="HTML", reply_markup=kb)
    await call.answer()


@dp.callback_query(F.data.startswith("debt:view:"))
async def debt_view(call: CallbackQuery, state: FSMContext):
    await state.clear()
    debt_id = int(call.data.split(":")[2])
    debt = get_debt(debt_id, uid(call))
    if not debt:
        await call.answer("Долг не найден.", show_alert=True)
        return
    _, amount, remaining, comment, is_closed, created_at = debt
    status = "✅ закрыт" if is_closed else f"остаток: <b>{remaining:g}</b> ₽"
    cmt = comment if comment and comment != "—" else "—"
    text = (f"💳 <b>Долг #{debt_id}</b>\n\n"
            f"Изначальная сумма: <b>{amount:g}</b> ₽\n"
            f"Статус: {status}\n💬 {cmt}\n📅 Создан: {created_at[:10]}\n")
    await call.message.edit_text(text, parse_mode="HTML",
                                 reply_markup=debt_actions_kb(debt_id))
    await call.answer()


@dp.callback_query(F.data.startswith("debt:history:"))
async def debt_history(call: CallbackQuery, state: FSMContext):
    debt_id = int(call.data.split(":")[2])
    debt = get_debt(debt_id, uid(call))
    if not debt:
        await call.answer("Долг не найден.", show_alert=True)
        return
    payments = get_debt_payments(debt_id, uid(call))
    _, amount, remaining, comment, is_closed, created_at = debt
    cmt = comment if comment and comment != "—" else "—"
    lines = [f"📜 <b>История по долгу #{debt_id}</b>\n",
             f"💬 {cmt}", f"📅 Создан: {created_at[:10]}",
             f"Сумма: <b>{amount:g}</b> ₽"]
    if is_closed:
        lines.append("Статус: ✅ закрыт")
    else:
        lines.append(f"Остаток: <b>{remaining:g}</b> ₽")
    lines.append("")
    if not payments:
        lines.append("Платежей пока не было.")
    else:
        paid_total = sum(p[0] for p in payments)
        lines.append(f"💸 <b>Платежи ({len(payments)} шт., всего {paid_total:g} ₽):</b>")
        for amt, dt in payments:
            lines.append(f"  • {dt[:16]} — <b>{amt:g}</b> ₽")
        lines.append("")
        if amount > 0:
            percent = min(paid_total / amount * 100, 100)
            lines.append(f"📊 Прогресс: {percent:.1f}%")
    lines.append("")
    lines.append("<i>Отменить последний платёж можно в меню долга.</i>")
    text = "\n".join(lines)
    if len(text) > 4000:
        text = text[:4000] + "\n… (обрезано)"
    await call.message.edit_text(
        text, parse_mode="HTML",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="⬅️ К долгу",
                                  callback_data=f"debt:view:{debt_id}")],
            [InlineKeyboardButton(text="🏠 К списку",
                                  callback_data="debt:list")]]))
    await call.answer()


@dp.callback_query(F.data.startswith("debt:edit_amount:"))
async def debt_edit_amount_start(call: CallbackQuery, state: FSMContext):
    debt_id = int(call.data.split(":")[2])
    debt = get_debt(debt_id, uid(call))
    if not debt:
        await call.answer("Долг не найден.", show_alert=True)
        return
    await state.set_state(DebtEditFlow.amount)
    await state.update_data(debt_id=debt_id)
    _, amount, remaining, *_ = debt
    await call.message.edit_text(
        f"✏️ <b>Изменение суммы долга #{debt_id}</b>\n\n"
        f"Текущая сумма: <b>{amount:g}</b> ₽\n"
        f"Остаток: <b>{remaining:g}</b> ₽\n\nВведите новую сумму:",
        parse_mode="HTML")
    await call.answer()


@dp.message(DebtEditFlow.amount)
async def debt_edit_amount_input(message: Message, state: FSMContext):
    try:
        new_amount = float(message.text.replace(",", ".").strip())
        if new_amount <= 0:
            raise ValueError
    except ValueError:
        await message.answer("❗ Введите корректное положительное число.")
        return
    data = await state.get_data()
    debt_id = data["debt_id"]
    result = update_debt(debt_id, uid(message), new_amount=new_amount)
    await state.clear()
    if not result:
        await message.answer("Долг не найден.",
                             reply_markup=_menu_for(uid(message)))
        return
    new_amount_db, new_remaining = result
    status = ("✅ долг закрыт (остаток 0)" if new_remaining == 0
              else f"остаток: <b>{new_remaining:g}</b> ₽")
    await message.answer(
        f"✅ Сумма долга #{debt_id} обновлена.\n\n"
        f"Новая сумма: <b>{new_amount_db:g}</b> ₽\nСтатус: {status}",
        parse_mode="HTML", reply_markup=_menu_for(uid(message)))


@dp.callback_query(F.data.startswith("debt:edit_comment:"))
async def debt_edit_comment_start(call: CallbackQuery, state: FSMContext):
    debt_id = int(call.data.split(":")[2])
    debt = get_debt(debt_id, uid(call))
    if not debt:
        await call.answer("Долг не найден.", show_alert=True)
        return
    await state.set_state(DebtEditFlow.comment)
    await state.update_data(debt_id=debt_id)
    _, _, _, comment, *_ = debt
    cmt = comment if comment and comment != "—" else "—"
    await call.message.edit_text(
        f"✏️ <b>Изменение комментария долга #{debt_id}</b>\n\n"
        f"Текущий: <i>{cmt}</i>\n\n"
        f"Введите новый комментарий (или нажмите «Убрать»):",
        parse_mode="HTML",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="🗑 Убрать комментарий",
                                  callback_data=f"debt:clear_comment:{debt_id}")],
            [InlineKeyboardButton(text="⬅️ Отмена",
                                  callback_data=f"debt:view:{debt_id}")]]))
    await call.answer()


@dp.message(DebtEditFlow.comment)
async def debt_edit_comment_input(message: Message, state: FSMContext):
    data = await state.get_data()
    debt_id = data["debt_id"]
    new_comment = message.text.strip()
    result = update_debt(debt_id, uid(message), new_comment=new_comment)
    await state.clear()
    if not result:
        await message.answer("Долг не найден.",
                             reply_markup=_menu_for(uid(message)))
        return
    await message.answer(
        f"✅ Комментарий долга #{debt_id} обновлён:\n💬 <i>{new_comment}</i>",
        parse_mode="HTML", reply_markup=_menu_for(uid(message)))


@dp.callback_query(F.data.startswith("debt:clear_comment:"))
async def debt_clear_comment(call: CallbackQuery, state: FSMContext):
    debt_id = int(call.data.split(":")[2])
    result = update_debt(debt_id, uid(call), new_comment="—")
    await state.clear()
    if not result:
        await call.answer("Долг не найден.", show_alert=True)
        return
    await call.message.edit_text(
        f"✅ Комментарий долга #{debt_id} убран.",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="⬅️ К долгу",
                                  callback_data=f"debt:view:{debt_id}")]]))
    await call.answer()


@dp.callback_query(F.data.startswith("debt:rollback:"))
async def debt_rollback(call: CallbackQuery, state: FSMContext):
    debt_id = int(call.data.split(":")[2])
    last = get_last_payment(debt_id, uid(call))
    if not last:
        await call.answer("По этому долгу ещё нет платежей.", show_alert=True)
        return
    _, amount, created_at = last
    await call.message.edit_text(
        f"↩️ <b>Отмена последнего платежа</b>\n\n"
        f"Долг #{debt_id}\nПлатёж: <b>{amount:g}</b> ₽\n"
        f"Дата: {created_at[:16]}\n\n"
        f"Отменить этот платёж? Остаток будет восстановлен.",
        parse_mode="HTML",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="✅ Да, отменить",
                                  callback_data=f"debt:rollback_confirm:{debt_id}")],
            [InlineKeyboardButton(text="❌ Нет",
                                  callback_data=f"debt:view:{debt_id}")]]))
    await call.answer()


@dp.callback_query(F.data.startswith("debt:rollback_confirm:"))
async def debt_rollback_confirm(call: CallbackQuery, state: FSMContext):
    debt_id = int(call.data.split(":")[2])
    result = rollback_last_payment(debt_id, uid(call))
    if not result:
        await call.answer("Не удалось отменить.", show_alert=True)
        return
    removed_amount, new_remaining = result
    status_line = ("✅ Долг закрыт (остаток 0)" if new_remaining == 0
                   else f"Остаток: <b>{new_remaining:g}</b> ₽")
    await call.message.edit_text(
        f"↩️ Платёж <b>{removed_amount:g}</b> ₽ отменён.\n"
        f"Долг #{debt_id}\n{status_line}",
        parse_mode="HTML",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="⬅️ К долгу",
                                  callback_data=f"debt:view:{debt_id}")],
            [InlineKeyboardButton(text="🏠 К списку",
                                  callback_data="debt:list")]]))
    await call.answer()


# ============================================================
#                       ИНВЕСТИЦИИ
# ============================================================
@dp.message(F.text == "📈 Инвестиции")
async def invest_menu(message: Message, state: FSMContext):
    await state.clear()
    total_in, total_out, balance = get_investment_totals(uid(message))
    realized = get_total_realized_profit(uid(message))
    pv = get_portfolio_value(uid(message))
    lines = ["📈 <b>Инвестиции</b>\n",
             f"📥 Всего вложено: <b>{total_in:g}</b> ₽",
             f"📤 Всего выведено: <b>{total_out:g}</b> ₽",
             f"📊 Баланс: <b>{balance:g}</b> ₽",
             f"💰 Зафиксировано прибыли: <b>{realized:+g}</b> ₽"]
    if pv:
        value, _ = pv
        pnl = value - balance
        arrow = "🔺" if pnl > 0 else ("🔻" if pnl < 0 else "➖")
        pct = (pnl / balance * 100) if balance else 0
        lines.append("")
        lines.append(f"💼 Текущая стоимость: <b>{value:g}</b> ₽")
        lines.append(f"{arrow} P&L портфеля: <b>{pnl:+g}</b> ₽ ({pct:+.1f}%)")
    await message.answer("\n".join(lines), parse_mode="HTML",
                         reply_markup=invest_menu_kb())


@dp.callback_query(F.data.in_({"invest:deposit", "invest:withdraw"}))
async def invest_start(call: CallbackQuery, state: FSMContext):
    inv_type = "deposit" if call.data == "invest:deposit" else "withdraw"
    await state.set_state(InvestFlow.amount)
    await state.update_data(inv_type=inv_type)
    title = "пополнения" if inv_type == "deposit" else "вывода"
    await call.message.edit_text(f"Введите сумму {title}:")
    await call.answer()


@dp.message(InvestFlow.amount)
async def invest_amount(message: Message, state: FSMContext):
    try:
        amount = float(message.text.replace(",", ".").strip())
        if amount <= 0:
            raise ValueError
    except ValueError:
        await message.answer("❗ Введите корректное положительное число.")
        return
    await state.update_data(amount=amount)
    await state.set_state(InvestFlow.category)
    cats = get_categories(uid(message), "invest")
    await message.answer("Выберите категорию:",
                         reply_markup=category_kb(cats))


@dp.callback_query(InvestFlow.category, F.data.startswith("cat:"))
async def invest_category(call: CallbackQuery, state: FSMContext):
    category = call.data.split(":", 1)[1]
    await state.update_data(category=category)
    await state.set_state(InvestFlow.comment)
    await call.message.edit_text(
        f"Категория: {category}\n\n"
        f"Напишите комментарий или нажмите «Без комментария»:",
        reply_markup=skip_invest_comment_kb())
    await call.answer()


@dp.message(InvestFlow.comment)
async def invest_comment(message: Message, state: FSMContext):
    await _save_investment(message, state, message.text.strip())


@dp.callback_query(InvestFlow.comment, F.data == "invest:skip_comment")
async def invest_skip_comment(call: CallbackQuery, state: FSMContext):
    await _save_investment(call, state, "—")


async def _save_investment(event, state: FSMContext, comment: str):
    data = await state.get_data()
    add_investment(user_id=uid(event), inv_type=data["inv_type"],
                   amount=data["amount"], category=data["category"],
                   comment=comment)
    await state.clear()
    kind = "📥 Пополнение" if data["inv_type"] == "deposit" else "📤 Вывод"
    text = f"✅ {kind} <b>{data['amount']:g}</b> ₽ записано ({data['category']})."
    if isinstance(event, CallbackQuery):
        await event.message.edit_text(text, parse_mode="HTML")
        await event.message.answer("Главное меню:", reply_markup=_menu_for(uid(event)))
        await event.answer()
    else:
        await event.answer(text, parse_mode="HTML",
                           reply_markup=_menu_for(uid(event)))


@dp.callback_query(F.data == "invest:summary")
async def invest_summary(call: CallbackQuery, state: FSMContext):
    user_id = uid(call)
    total_in, total_out, balance = get_investment_totals(user_id)
    by_cat = get_investments_by_category(user_id)
    realized = get_total_realized_profit(user_id)
    realized_by_cat = get_profits_by_category(user_id)
    pv = get_portfolio_value(user_id)
    lines = ["📊 <b>Сводка по инвестициям</b>\n",
             f"📥 Вложено: <b>{total_in:g}</b> ₽",
             f"📤 Выведено: <b>{total_out:g}</b> ₽",
             f"📊 В портфеле (нетто): <b>{balance:g}</b> ₽\n"]
    if by_cat:
        lines.append("💼 <b>По категориям (вложено):</b>")
        for cat, s in by_cat:
            pct = (s / total_in * 100) if total_in else 0
            lines.append(f"  • {cat}: <b>{s:g}</b> ₽ ({pct:.1f}%)")
        lines.append("")
    unrealized = None
    if pv:
        value, _ = pv
        unrealized = value - balance
    lines.append("💵 <b>Прибыль:</b>")
    lines.append(f"  💰 Зафиксировано: <b>{realized:+g}</b> ₽")
    if unrealized is not None:
        lines.append(f"  💼 В портфеле: <b>{unrealized:+g}</b> ₽")
        total_pnl = realized + unrealized
        arrow = "🔺" if total_pnl > 0 else ("🔻" if total_pnl < 0 else "➖")
        lines.append(f"  {arrow} <b>Итого: {total_pnl:+g} ₽</b>")
    else:
        lines.append("  💼 В портфеле: <i>укажите текущую стоимость</i>")
    if realized_by_cat:
        lines.append("")
        lines.append("💰 <b>Зафиксированная прибыль по категориям:</b>")
        for cat, s in realized_by_cat:
            lines.append(f"  • {cat}: <b>{s:+g}</b> ₽")
    await call.message.edit_text("\n".join(lines), parse_mode="HTML",
                                 reply_markup=invest_menu_kb())
    await call.answer()


@dp.callback_query(F.data == "invest:list")
async def invest_list(call: CallbackQuery, state: FSMContext):
    rows = get_investments(uid(call), limit=30)
    if not rows:
        await call.message.edit_text("Пока нет операций по инвестициям.",
                                     reply_markup=invest_menu_kb())
        await call.answer()
        return
    lines = ["📋 <b>Последние операции:</b>\n"]
    for inv_id, inv_type, amount, category, comment, created_at in rows:
        emoji = "📥" if inv_type == "deposit" else "📤"
        cmt = comment if comment and comment != "—" else ""
        tail = f" — {cmt}" if cmt else ""
        lines.append(f"{emoji} <b>{amount:g}</b> ₽ | {category} | {created_at[:16]}{tail}")
    text = "\n".join(lines)
    if len(text) > 4000:
        text = text[:4000] + "\n… (обрезано)"
    await call.message.edit_text(text, parse_mode="HTML",
                                 reply_markup=invest_menu_kb())
    await call.answer()


@dp.callback_query(F.data == "invest:set_value")
async def invest_set_value(call: CallbackQuery, state: FSMContext):
    await state.set_state(InvestValueFlow.value)
    current = get_portfolio_value(uid(call))
    cur_line = f"\nТекущее значение: <b>{current[0]:g}</b> ₽" if current else ""
    await call.message.edit_text(
        f"Введите <b>текущую стоимость портфеля</b> в рублях.{cur_line}",
        parse_mode="HTML")
    await call.answer()


@dp.message(InvestValueFlow.value)
async def invest_set_value_input(message: Message, state: FSMContext):
    try:
        value = float(message.text.replace(",", ".").strip())
        if value < 0:
            raise ValueError
    except ValueError:
        await message.answer("❗ Введите корректное неотрицательное число.")
        return
    set_portfolio_value(uid(message), value)
    await state.clear()
    _, _, balance = get_investment_totals(uid(message))
    pnl = value - balance
    arrow = "🔺" if pnl > 0 else ("🔻" if pnl < 0 else "➖")
    pct = (pnl / balance * 100) if balance else 0
    await message.answer(
        f"✅ Текущая стоимость портфеля: <b>{value:g}</b> ₽\n"
        f"{arrow} P&L: <b>{pnl:+g}</b> ₽ ({pct:+.1f}%)",
        parse_mode="HTML", reply_markup=_menu_for(uid(message)))


# ============================================================
#                       ПРИБЫЛИ
# ============================================================
@dp.callback_query(F.data == "profit:add")
async def profit_start(call: CallbackQuery, state: FSMContext):
    await state.set_state(ProfitFlow.amount)
    await call.message.edit_text(
        "Введите сумму прибыли.\n\n"
        "💡 Если получили <b>убыток</b> — введите со знаком минус "
        "(например, <code>-5000</code>).",
        parse_mode="HTML")
    await call.answer()


@dp.message(ProfitFlow.amount)
async def profit_amount(message: Message, state: FSMContext):
    raw = message.text.replace(",", ".").strip()
    try:
        amount = float(raw)
        if amount == 0:
            raise ValueError
    except ValueError:
        await message.answer("❗ Введите число, отличное от нуля.")
        return
    await state.update_data(amount=amount)
    await state.set_state(ProfitFlow.category)
    cats = get_categories(uid(message), "invest")
    await message.answer("Выберите категорию:",
                         reply_markup=category_kb(cats))


@dp.callback_query(ProfitFlow.category, F.data.startswith("cat:"))
async def profit_category(call: CallbackQuery, state: FSMContext):
    category = call.data.split(":", 1)[1]
    await state.update_data(category=category)
    await state.set_state(ProfitFlow.source)
    sources = get_categories(uid(call), "profit_source")
    await call.message.edit_text(
        f"Категория: {category}\n\nКак получена прибыль?",
        reply_markup=profit_source_kb(sources))
    await call.answer()


@dp.callback_query(ProfitFlow.source, F.data.startswith("psrc:"))
async def profit_source(call: CallbackQuery, state: FSMContext):
    source = call.data.split(":", 1)[1]
    await state.update_data(source=source)
    await state.set_state(ProfitFlow.comment)
    await call.message.edit_text(
        f"Источник: {source}\n\n"
        f"Напишите комментарий или нажмите «Без комментария»:",
        reply_markup=skip_profit_comment_kb())
    await call.answer()


@dp.message(ProfitFlow.comment)
async def profit_comment(message: Message, state: FSMContext):
    await _save_profit(message, state, message.text.strip())


@dp.callback_query(ProfitFlow.comment, F.data == "profit:skip_comment")
async def profit_skip_comment(call: CallbackQuery, state: FSMContext):
    await _save_profit(call, state, "—")


async def _save_profit(event, state: FSMContext, comment: str):
    data = await state.get_data()
    add_profit(user_id=uid(event), amount=data["amount"],
               category=data["category"], source=data["source"],
               comment=comment)
    await state.clear()
    amount = data["amount"]
    kind = "💰 Прибыль" if amount > 0 else "📉 Убыток"
    text = (f"✅ {kind} <b>{amount:+g}</b> ₽ зафиксирован.\n"
            f"Категория: {data['category']} | {data['source']}")
    if isinstance(event, CallbackQuery):
        await event.message.edit_text(text, parse_mode="HTML")
        await event.message.answer("Главное меню:", reply_markup=_menu_for(uid(event)))
        await event.answer()
    else:
        await event.answer(text, parse_mode="HTML",
                           reply_markup=_menu_for(uid(event)))


@dp.callback_query(F.data == "profit:list")
async def profit_list(call: CallbackQuery, state: FSMContext):
    rows = get_profits(uid(call), limit=50)
    if not rows:
        await call.message.edit_text("Пока нет зафиксированных прибылей.",
                                     reply_markup=invest_menu_kb())
        await call.answer()
        return
    total = sum(r[1] for r in rows)
    lines = [f"📜 <b>История прибылей</b>", f"<i>Всего: {total:+g} ₽</i>\n"]
    for pid, amount, category, source, comment, created_at in rows:
        emoji = "💰" if amount > 0 else "📉"
        cmt = comment if comment and comment != "—" else ""
        tail = f" — {cmt}" if cmt else ""
        lines.append(f"{emoji} <b>{amount:+g}</b> ₽ | {category} | {source} "
                     f"| {created_at[:16]}{tail}")
    text = "\n".join(lines)
    if len(text) > 4000:
        text = text[:4000] + "\n… (обрезано)"
    await call.message.edit_text(text, parse_mode="HTML",
                                 reply_markup=invest_menu_kb())
    await call.answer()


# ============================================================
#                 НАСТРОЙКИ ОТЧЁТОВ / НАПОМИНАНИЙ
# ============================================================
@dp.message(F.text == "🗓 Еженедельный отчёт")
async def weekly_menu(message: Message, state: FSMContext):
    await state.clear()
    s = get_settings(uid(message))
    await message.answer(
        "Настройки еженедельного отчёта.\n"
        "Отчёт приходит каждое <b>воскресенье</b> в указанное время.",
        parse_mode="HTML",
        reply_markup=weekly_menu_kb(bool(s["weekly_enabled"]), s["weekly_time"],
                                    s["tz_offset"], bool(s["weekly_send_excel"])))


@dp.callback_query(F.data == "noop")
async def noop(call: CallbackQuery):
    await call.answer()


@dp.callback_query(F.data == "weekly:set_time")
async def weekly_set_time(call: CallbackQuery, state: FSMContext):
    await state.set_state(WeeklyReportFlow.time_input)
    await call.message.edit_text("Введите время в формате <b>HH:MM</b>:",
                                 parse_mode="HTML")
    await call.answer()


@dp.message(WeeklyReportFlow.time_input)
async def weekly_time_input(message: Message, state: FSMContext):
    raw = message.text.strip()
    try:
        hh, mm = raw.split(":")
        hh, mm = int(hh), int(mm)
        if not (0 <= hh < 24 and 0 <= mm < 60):
            raise ValueError
    except Exception:
        await message.answer("❗ Формат: 20:00")
        return
    time_str = f"{hh:02d}:{mm:02d}"
    save_settings(uid(message), weekly_enabled=1, weekly_time=time_str)
    await state.clear()
    s = get_settings(uid(message))
    await message.answer(
        f"✅ Отчёт будет приходить в воскресенье в <b>{time_str}</b> "
        f"(UTC{s['tz_offset']:+d}).",
        parse_mode="HTML", reply_markup=_menu_for(uid(message)))


@dp.callback_query(F.data == "weekly:set_tz")
async def weekly_set_tz(call: CallbackQuery, state: FSMContext):
    await state.set_state(WeeklyReportFlow.tz_input)
    await call.message.edit_text("Введите смещение от UTC (-12…14):")
    await call.answer()


@dp.message(WeeklyReportFlow.tz_input)
async def weekly_tz_input(message: Message, state: FSMContext):
    try:
        tz = int(message.text.strip())
        if not (-12 <= tz <= 14):
            raise ValueError
    except ValueError:
        await message.answer("❗ Число от -12 до 14.")
        return
    save_settings(uid(message), tz_offset=tz)
    await state.clear()
    await message.answer(f"✅ Часовой пояс: UTC{tz:+d}.",
                         reply_markup=_menu_for(uid(message)))


@dp.callback_query(F.data == "weekly:toggle_excel")
async def weekly_toggle_excel(call: CallbackQuery):
    s = get_settings(uid(call))
    new_val = 0 if s["weekly_send_excel"] else 1
    save_settings(uid(call), weekly_send_excel=new_val)
    s = get_settings(uid(call))
    await call.message.edit_reply_markup(reply_markup=weekly_menu_kb(
        bool(s["weekly_enabled"]), s["weekly_time"], s["tz_offset"],
        bool(s["weekly_send_excel"])))
    await call.answer("Автоотправка Excel " + ("включена" if new_val else "выключена"))


@dp.callback_query(F.data == "weekly:disable")
async def weekly_disable(call: CallbackQuery):
    save_settings(uid(call), weekly_enabled=0)
    await call.message.edit_text("🔕 Еженедельный отчёт выключен.")
    await call.answer()


@dp.message(F.text == "⏰ Ежедневный отчёт")
async def daily_menu(message: Message, state: FSMContext):
    await state.clear()
    s = get_settings(uid(message))
    await message.answer(
        "Настройки ежедневного отчёта.\n"
        "Отчёт приходит каждый день в указанное время.",
        reply_markup=daily_menu_kb(bool(s["daily_enabled"]), s["daily_time"],
                                   s["tz_offset"], bool(s["daily_send_excel"])))


@dp.callback_query(F.data == "daily:set_time")
async def daily_set_time(call: CallbackQuery, state: FSMContext):
    await state.set_state(DailyReportFlow.time_input)
    await call.message.edit_text("Введите время в формате <b>HH:MM</b>:",
                                 parse_mode="HTML")
    await call.answer()


@dp.message(DailyReportFlow.time_input)
async def daily_time_input(message: Message, state: FSMContext):
    raw = message.text.strip()
    try:
        hh, mm = raw.split(":")
        hh, mm = int(hh), int(mm)
        if not (0 <= hh < 24 and 0 <= mm < 60):
            raise ValueError
    except Exception:
        await message.answer("❗ Формат: 21:30")
        return
    time_str = f"{hh:02d}:{mm:02d}"
    save_settings(uid(message), daily_enabled=1, daily_time=time_str)
    await state.clear()
    await message.answer(f"✅ Ежедневный отчёт будет приходить в {time_str}.",
                         reply_markup=_menu_for(uid(message)))


@dp.callback_query(F.data == "daily:set_tz")
async def daily_set_tz(call: CallbackQuery, state: FSMContext):
    await state.set_state(DailyReportFlow.tz_input)
    await call.message.edit_text("Введите смещение от UTC (-12…14):")
    await call.answer()


@dp.message(DailyReportFlow.tz_input)
async def daily_tz_input(message: Message, state: FSMContext):
    try:
        tz = int(message.text.strip())
        if not (-12 <= tz <= 14):
            raise ValueError
    except ValueError:
        await message.answer("❗ Число от -12 до 14.")
        return
    save_settings(uid(message), tz_offset=tz)
    await state.clear()
    await message.answer(f"✅ Часовой пояс: UTC{tz:+d}.",
                         reply_markup=_menu_for(uid(message)))


@dp.callback_query(F.data == "daily:toggle_excel")
async def daily_toggle_excel(call: CallbackQuery):
    s = get_settings(uid(call))
    new_val = 0 if s["daily_send_excel"] else 1
    save_settings(uid(call), daily_send_excel=new_val)
    s = get_settings(uid(call))
    await call.message.edit_reply_markup(reply_markup=daily_menu_kb(
        bool(s["daily_enabled"]), s["daily_time"], s["tz_offset"],
        bool(s["daily_send_excel"])))
    await call.answer("Автоотправка Excel " + ("включена" if new_val else "выключена"))


@dp.callback_query(F.data == "daily:disable")
async def daily_disable(call: CallbackQuery):
    save_settings(uid(call), daily_enabled=0)
    await call.message.edit_text("🔕 Ежедневный отчёт выключен.")
    await call.answer()


@dp.message(F.text == "🔔 Напоминания")
async def reminder_menu(message: Message, state: FSMContext):
    await state.clear()
    s = get_settings(uid(message))
    await message.answer(
        "🔔 <b>Напоминания</b>\n\n"
        "Бот будет напоминать записать траты с указанным интервалом. "
        "В «тихие часы» не беспокоит.",
        parse_mode="HTML", reply_markup=reminder_menu_kb(s))


@dp.callback_query(F.data == "reminder:enable")
async def reminder_enable(call: CallbackQuery):
    save_settings(uid(call), reminder_enabled=1)
    s = get_settings(uid(call))
    await call.message.edit_reply_markup(reply_markup=reminder_menu_kb(s))
    await call.answer("Напоминания включены")


@dp.callback_query(F.data == "reminder:disable")
async def reminder_disable(call: CallbackQuery):
    save_settings(uid(call), reminder_enabled=0)
    s = get_settings(uid(call))
    await call.message.edit_reply_markup(reply_markup=reminder_menu_kb(s))
    await call.answer("Напоминания выключены")


@dp.callback_query(F.data == "reminder:interval")
async def reminder_set_interval(call: CallbackQuery, state: FSMContext):
    await state.set_state(ReminderFlow.interval)
    await call.message.edit_text("Введите интервал напоминаний в часах (1–24):")
    await call.answer()


@dp.message(ReminderFlow.interval)
async def reminder_interval_input(message: Message, state: FSMContext):
    try:
        v = int(message.text.strip())
        if not (1 <= v <= 24):
            raise ValueError
    except ValueError:
        await message.answer("❗ Число от 1 до 24.")
        return
    save_settings(uid(message), reminder_interval_hours=v, reminder_enabled=1)
    await state.clear()
    await message.answer(f"✅ Интервал: каждые {v} ч.",
                         reply_markup=_menu_for(uid(message)))


@dp.callback_query(F.data == "reminder:silent_from")
async def reminder_silent_from(call: CallbackQuery, state: FSMContext):
    await state.set_state(ReminderFlow.silent_from)
    await call.message.edit_text("С какого часа не беспокоить? (0–23)")
    await call.answer()


@dp.message(ReminderFlow.silent_from)
async def reminder_silent_from_input(message: Message, state: FSMContext):
    try:
        v = int(message.text.strip())
        if not (0 <= v <= 23):
            raise ValueError
    except ValueError:
        await message.answer("❗ Число от 0 до 23.")
        return
    save_settings(uid(message), reminder_silent_from=v)
    await state.set_state(ReminderFlow.silent_to)
    await message.answer("До какого часа? (0–23)")


@dp.message(ReminderFlow.silent_to)
async def reminder_silent_to_input(message: Message, state: FSMContext):
    try:
        v = int(message.text.strip())
        if not (0 <= v <= 23):
            raise ValueError
    except ValueError:
        await message.answer("❗ Число от 0 до 23.")
        return
    save_settings(uid(message), reminder_silent_to=v)
    await state.clear()
    await message.answer("✅ Тихие часы сохранены.",
                         reply_markup=_menu_for(uid(message)))


# ============================================================
#                  ЭКСПОРТ В EXCEL — МЕНЮ
# ============================================================
@dp.message(F.text == "📤 Экспорт в Excel")
async def excel_menu(message: Message, state: FSMContext):
    await state.clear()
    s = get_settings(uid(message))
    await message.answer("📤 <b>Экспорт в Excel</b>\n\nВыберите действие:",
                         parse_mode="HTML", reply_markup=excel_menu_kb(s))


@dp.callback_query(F.data == "xls:menu")
async def excel_menu_cb(call: CallbackQuery, state: FSMContext):
    await state.clear()
    s = get_settings(uid(call))
    await call.message.edit_text("📤 <b>Экспорт в Excel</b>\n\nВыберите действие:",
                                 parse_mode="HTML", reply_markup=excel_menu_kb(s))
    await call.answer()


@dp.callback_query(F.data.startswith("xls:toggle:"))
async def excel_toggle(call: CallbackQuery):
    key = call.data.split(":")[2]
    field_map = {"ops": "excel_include_ops", "debts": "excel_include_debts",
                 "payments": "excel_include_payments",
                 "invest": "excel_include_invest",
                 "profits": "excel_include_profits"}
    field = field_map.get(key)
    if not field:
        await call.answer("Неизвестный параметр.")
        return
    s = get_settings(uid(call))
    save_settings(uid(call), **{field: 0 if s[field] else 1})
    s = get_settings(uid(call))
    await call.message.edit_reply_markup(reply_markup=excel_menu_kb(s))
    await call.answer("Изменено")


@dp.callback_query(F.data.startswith("xls:auto:toggle:"))
async def excel_auto_toggle(call: CallbackQuery):
    key = call.data.split(":")[3]
    field = {"debt_close": "excel_on_debt_close",
             "big_op": "excel_on_big_op"}.get(key)
    if not field:
        await call.answer("Неизвестный параметр.")
        return
    s = get_settings(uid(call))
    save_settings(uid(call), **{field: 0 if s[field] else 1})
    s = get_settings(uid(call))
    await call.message.edit_reply_markup(reply_markup=excel_menu_kb(s))
    await call.answer("Изменено")


@dp.callback_query(F.data == "xls:auto:big_op_threshold")
async def excel_big_op_threshold(call: CallbackQuery, state: FSMContext):
    await state.set_state(ExcelAutoFlow.threshold)
    await call.message.edit_text("Введите порог крупной операции (например, 10000):")
    await call.answer()


@dp.message(ExcelAutoFlow.threshold)
async def excel_big_op_threshold_input(message: Message, state: FSMContext):
    try:
        v = float(message.text.replace(",", ".").strip())
        if v <= 0:
            raise ValueError
    except ValueError:
        await message.answer("❗ Введите положительное число.")
        return
    save_settings(uid(message), excel_big_op_threshold=v)
    await state.clear()
    await message.answer(f"✅ Порог: {v:g} ₽.",
                         reply_markup=_menu_for(uid(message)))


@dp.callback_query(F.data == "xls:pick_month")
async def excel_pick_month(call: CallbackQuery, state: FSMContext):
    months = get_months(uid(call))
    if not months:
        await call.answer("Нет данных.", show_alert=True)
        return
    await call.message.edit_text("Выберите месяц:",
                                 reply_markup=months_kb(months, "xls"))


@dp.callback_query(F.data.startswith("xls:"))
async def excel_export_month(call: CallbackQuery, state: FSMContext):
    if (call.data.startswith("xls:toggle") or
            call.data.startswith("xls:auto") or
            call.data.startswith("xls:hist") or
            call.data in ("xls:pick_month", "xls:menu", "xls:range")):
        return
    month = call.data.split(":", 1)[1]
    if "-" not in month:
        return
    await call.message.edit_text("⏳ Формирую файл...")
    y, mo = map(int, month.split("-"))
    start = f"{y:04d}-{mo:02d}-01"
    next_month = (f"{y+1:04d}-01-01" if mo == 12 else f"{y:04d}-{mo+1:02d}-01")
    end_dt = datetime.strptime(next_month, "%Y-%m-%d") - timedelta(days=1)
    end = end_dt.strftime("%Y-%m-%d")
    try:
        await _send_excel(uid(call), start, end,
                          title=f"📤 Экспорт за {month}",
                          reason="manual", filename=f"finance_{month}.xlsx")
        await call.answer()
    except Exception as e:
        logging.exception("Ошибка экспорта")
        await call.answer(f"Ошибка: {e}", show_alert=True)


@dp.callback_query(F.data == "xls:range")
async def excel_range_start(call: CallbackQuery, state: FSMContext):
    await state.set_state(ExcelRangeFlow.start_date)
    await state.update_data(csv_mode=False)
    await call.message.edit_text(
        "Введите <b>начальную дату</b> в формате <code>ГГГГ-ММ-ДД</code>:",
        parse_mode="HTML")
    await call.answer()


@dp.message(ExcelRangeFlow.start_date)
async def excel_range_start_input(message: Message, state: FSMContext):
    raw = message.text.strip()
    try:
        datetime.strptime(raw, "%Y-%m-%d")
    except ValueError:
        await message.answer("❗ Формат: ГГГГ-ММ-ДД.")
        return
    await state.update_data(start=raw)
    await state.set_state(ExcelRangeFlow.end_date)
    await message.answer("Введите <b>конечную дату</b> (ГГГГ-ММ-ДД):",
                         parse_mode="HTML")


@dp.message(ExcelRangeFlow.end_date)
async def excel_range_end_input(message: Message, state: FSMContext):
    raw = message.text.strip()
    try:
        dt_end = datetime.strptime(raw, "%Y-%m-%d")
    except ValueError:
        await message.answer("❗ Формат: ГГГГ-ММ-ДД.")
        return
    data = await state.get_data()
    start = data["start"]
    csv_mode = data.get("csv_mode", False)
    if dt_end < datetime.strptime(start, "%Y-%m-%d"):
        await message.answer("❗ Конечная дата раньше начальной.")
        return
    await state.clear()
    status_msg = await message.answer("⏳ Формирую файл...")
    try:
        if csv_mode:
            rows = get_export_history_by_period(uid(message), start, raw)
            if not rows:
                await message.answer("За этот период нет записей.")
                return
            lines = ["ID;Файл;Причина;Период с;Период по;Размер, КБ;Дата отправки"]
            for exp_id, filename, reason, ps, pe, size, created_at in rows:
                size_kb = f"{size/1024:.1f}" if size else "0"
                safe_filename = filename.replace(";", ",")
                lines.append(f"{exp_id};{safe_filename};"
                             f"{REASON_LABELS.get(reason, reason)};"
                             f"{ps or ''};{pe or ''};{size_kb};{created_at}")
            csv_data = "\n".join(lines).encode("utf-8-sig")
            file = BufferedInputFile(csv_data,
                                     filename=f"export_history_{start}_to_{raw}.csv")
            await message.answer_document(file,
                                          caption=f"📥 История за {start} — {raw}")
        else:
            await _send_excel(uid(message), start, raw,
                              title=f"📤 Экспорт за {start} — {raw}",
                              reason="range",
                              filename=f"finance_{start}_to_{raw}.xlsx")
        try:
            await status_msg.delete()
        except Exception:
            pass
    except Exception as e:
        logging.exception("Ошибка экспорта")
        await message.answer(f"Ошибка: {e}", reply_markup=_menu_for(uid(message)))


@dp.callback_query(F.data == "auto:xls:menu")
async def autoexport_menu(call: CallbackQuery, state: FSMContext):
    await state.clear()
    s = get_settings(uid(call))
    await call.message.edit_text(
        "🗓 <b>Автоэкспорт Excel по расписанию</b>",
        parse_mode="HTML", reply_markup=autoexport_menu_kb(s))
    await call.answer()


@dp.callback_query(F.data == "auto:xls:enable")
async def autoexport_enable(call: CallbackQuery):
    save_settings(uid(call), autoexport_enabled=1)
    s = get_settings(uid(call))
    await call.message.edit_reply_markup(reply_markup=autoexport_menu_kb(s))
    await call.answer("Автоэкспорт включён")


@dp.callback_query(F.data == "auto:xls:disable")
async def autoexport_disable(call: CallbackQuery):
    save_settings(uid(call), autoexport_enabled=0)
    s = get_settings(uid(call))
    await call.message.edit_reply_markup(reply_markup=autoexport_menu_kb(s))
    await call.answer("Автоэкспорт выключен")


@dp.callback_query(F.data == "auto:xls:weekday")
async def autoexport_weekday(call: CallbackQuery, state: FSMContext):
    await state.set_state(AutoExportFlow.weekday)
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text=wd, callback_data=f"auto:xls:wd:{i}")]
        for i, wd in enumerate(WEEKDAYS_RU)])
    await call.message.edit_text("Выберите день недели:", reply_markup=kb)
    await call.answer()


@dp.callback_query(F.data.startswith("auto:xls:wd:"))
async def autoexport_weekday_set(call: CallbackQuery, state: FSMContext):
    idx = int(call.data.split(":")[3])
    save_settings(uid(call), autoexport_weekday=idx)
    await state.clear()
    s = get_settings(uid(call))
    await call.message.edit_text("🗓 <b>Автоэкспорт Excel</b>",
                                 parse_mode="HTML", reply_markup=autoexport_menu_kb(s))
    await call.answer("Сохранено")


@dp.callback_query(F.data == "auto:xls:time")
async def autoexport_time(call: CallbackQuery, state: FSMContext):
    await state.set_state(AutoExportFlow.time_input)
    await call.message.edit_text("Введите время (HH:MM):", parse_mode="HTML")
    await call.answer()


@dp.message(AutoExportFlow.time_input)
async def autoexport_time_input(message: Message, state: FSMContext):
    raw = message.text.strip()
    try:
        hh, mm = raw.split(":")
        hh, mm = int(hh), int(mm)
        if not (0 <= hh < 24 and 0 <= mm < 60):
            raise ValueError
    except Exception:
        await message.answer("❗ Формат: 20:00")
        return
    save_settings(uid(message), autoexport_time=f"{hh:02d}:{mm:02d}")
    await state.clear()
    await message.answer(f"✅ Время: {hh:02d}:{mm:02d}.",
                         reply_markup=_menu_for(uid(message)))


@dp.callback_query(F.data == "auto:xls:period")
async def autoexport_period(call: CallbackQuery, state: FSMContext):
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="За неделю", callback_data="auto:xls:p:week")],
        [InlineKeyboardButton(text="За месяц", callback_data="auto:xls:p:month")],
        [InlineKeyboardButton(text="За всё время", callback_data="auto:xls:p:all")]])
    await call.message.edit_text("Какой период выгружать?", reply_markup=kb)
    await call.answer()


@dp.callback_query(F.data.startswith("auto:xls:p:"))
async def autoexport_period_set(call: CallbackQuery, state: FSMContext):
    period = call.data.split(":")[3]
    if period not in ("week", "month", "all"):
        await call.answer("Неверный период.")
        return
    save_settings(uid(call), autoexport_period=period)
    s = get_settings(uid(call))
    await call.message.edit_text("🗓 <b>Автоэкспорт Excel</b>",
                                 parse_mode="HTML", reply_markup=autoexport_menu_kb(s))
    await call.answer("Сохранено")


@dp.callback_query(F.data == "xls:history")
async def export_history_cb(call: CallbackQuery, state: FSMContext):
    await state.clear()
    await _render_history(call, reason_filter=None)


@dp.callback_query(F.data.startswith("xls:hist:f:"))
async def export_history_filter(call: CallbackQuery):
    val = call.data.split(":")[3]
    await _render_history(call, reason_filter=None if val == "all" else val)


async def _render_history(call: CallbackQuery, reason_filter):
    rows = get_export_history_filtered(uid(call), reason_filter=reason_filter, limit=20)
    header = "📜 <b>История экспортов</b>"
    if reason_filter:
        header += f" — {REASON_LABELS.get(reason_filter, reason_filter)}"
    if not rows:
        await call.message.edit_text(header + "\n\n<i>Пока ничего нет.</i>",
                                     parse_mode="HTML",
                                     reply_markup=history_filter_kb(reason_filter))
        await call.answer()
        return
    lines = [header + "\n"]
    for exp_id, filename, reason, ps, pe, size, created_at in rows:
        r_label = REASON_LABELS.get(reason, reason)
        kb_size = f"{size / 1024:.1f} КБ" if size else "—"
        period = ""
        if ps and pe:
            period = f" | {ps}" if ps == pe else f" | {ps} → {pe}"
        lines.append(f"<b>#{exp_id}</b> {r_label}\n"
                     f"  📅 {created_at[:16]}{period}\n"
                     f"  📎 {filename} ({kb_size})")
    text = "\n".join(lines)
    if len(text) > 4000:
        text = text[:4000] + "\n… (обрезано)"
    await call.message.edit_text(text, parse_mode="HTML",
                                 reply_markup=history_filter_kb(reason_filter))
    action_kb_rows = []
    for exp_id, filename, *_ in rows[:5]:
        action_kb_rows.append([
            InlineKeyboardButton(text=f"#{exp_id} 📎",
                                 callback_data=f"xls:hist:resend:{exp_id}"),
            InlineKeyboardButton(text=f"#{exp_id} 🗑",
                                 callback_data=f"xls:hist:del:{exp_id}")])
    if action_kb_rows:
        await call.message.answer(
            "🔧 Действия по последним записям:",
            reply_markup=InlineKeyboardMarkup(inline_keyboard=action_kb_rows))
    await call.answer()


@dp.callback_query(F.data.startswith("xls:hist:resend:"))
async def export_resend(call: CallbackQuery):
    exp_id = int(call.data.split(":")[3])
    row = get_export_by_id(exp_id, uid(call))
    if not row:
        await call.answer("Запись не найдена.", show_alert=True)
        return
    _, filename, filepath, *_ = row
    if not filepath or not os.path.exists(filepath):
        await call.answer("Файл недоступен на диске.", show_alert=True)
        return
    try:
        with open(filepath, "rb") as f:
            data = f.read()
        file = BufferedInputFile(data, filename=filename)
        await bot.send_document(uid(call), document=file,
                                caption=f"📎 Повторная отправка #{exp_id}")
        await call.answer("Файл отправлен")
    except Exception as e:
        await call.answer(f"Ошибка: {e}", show_alert=True)


@dp.callback_query(F.data.startswith("xls:hist:del:"))
async def export_delete(call: CallbackQuery):
    exp_id = int(call.data.split(":")[3])
    row = get_export_by_id(exp_id, uid(call))
    if not row:
        await call.answer("Запись не найдена.", show_alert=True)
        return
    _, filename, *_ = row
    await call.message.answer(
        f"🗑 Удалить запись #{exp_id}?\nФайл: {filename}",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="✅ Да, удалить",
                                  callback_data=f"xls:hist:del_confirm:{exp_id}")],
            [InlineKeyboardButton(text="❌ Отмена",
                                  callback_data="xls:history")]]))
    await call.answer()


@dp.callback_query(F.data.startswith("xls:hist:del_confirm:"))
async def export_delete_confirm(call: CallbackQuery):
    exp_id = int(call.data.split(":")[3])
    row = get_export_by_id(exp_id, uid(call))
    if not row:
        await call.answer("Запись не найдена.", show_alert=True)
        return
    _, filename, filepath, *_ = row
    if filepath and os.path.exists(filepath):
        try:
            os.remove(filepath)
        except Exception as e:
            logging.warning(f"Не удалось удалить {filepath}: {e}")
    delete_export(exp_id, uid(call))
    await call.message.edit_text(f"✅ Запись #{exp_id} удалена.")
    await call.answer()


@dp.callback_query(F.data == "xls:hist:csv")
async def export_history_csv(call: CallbackQuery, state: FSMContext):
    await state.set_state(ExcelRangeFlow.start_date)
    await state.update_data(csv_mode=True)
    await call.message.edit_text(
        "📥 <b>Экспорт истории в CSV</b>\n\n"
        "Введите <b>начальную дату</b> (ГГГГ-ММ-ДД):",
        parse_mode="HTML")
    await call.answer()


# ============================================================
#                     ПЛАНИРОВЩИКИ
# ============================================================
async def check_and_send_weekly():
    now_utc = datetime.now(timezone.utc)
    for user_id, time_str, tz_offset, send_excel in get_all_weekly_users():
        try:
            hh, mm = map(int, time_str.split(":"))
        except Exception:
            continue
        user_now = now_utc + timedelta(hours=tz_offset)
        if user_now.weekday() != 6:
            continue
        if user_now.hour != hh or user_now.minute != mm:
            continue
        start = (user_now - timedelta(days=user_now.weekday())).replace(
            hour=0, minute=0, second=0, microsecond=0)
        end = user_now
        prev_start = start - timedelta(days=7)
        prev_end = start - timedelta(seconds=1)
        ops = get_operations_for_range(user_id, start.strftime("%Y-%m-%d"),
                                       end.strftime("%Y-%m-%d"))
        prev_ops = get_operations_for_range(user_id, prev_start.strftime("%Y-%m-%d"),
                                            prev_end.strftime("%Y-%m-%d"))
        prev_exp, prev_inc = 0.0, 0.0
        for op_type, amount, *_ in prev_ops:
            if op_type == "expense":
                prev_exp += amount
            else:
                prev_inc += amount
        if not ops:
            text = (f"🗓 <b>Отчёт за неделю</b>\n"
                    f"<i>{start.strftime('%d.%m')} — {end.strftime('%d.%m.%Y')}</i>\n\n"
                    f"За эту неделю не было ни трат, ни зачислений.")
            chart = None
        else:
            agg = _aggregate(ops)
            text = _make_weekly_report(
                start, end, agg, ops=ops,
                prev_totals=(prev_exp, prev_inc),
                expense_order=get_categories(user_id, "expense"),
                income_order=get_categories(user_id, "income"))
            chart = _build_weekly_chart(start, end, ops)
        try:
            if chart:
                caption = text if len(text) <= 1024 else text[:1000] + "…"
                await bot.send_photo(user_id,
                                     photo=BufferedInputFile(chart, filename="week.png"),
                                     caption=caption, parse_mode="HTML")
                if len(text) > 1024:
                    await bot.send_message(user_id, text, parse_mode="HTML")
            else:
                await bot.send_message(user_id, text, parse_mode="HTML")
        except Exception as e:
            logging.warning(f"Не отправить отчёт {user_id}: {e}")
        if send_excel:
            try:
                await _send_excel(user_id, start.strftime("%Y-%m-%d"),
                                  end.strftime("%Y-%m-%d"),
                                  title=f"📎 Excel за неделю "
                                        f"{start.strftime('%d.%m')} — {end.strftime('%d.%m.%Y')}",
                                  reason="weekly",
                                  filename=f"finance_week_{start.strftime('%Y%m%d')}.xlsx")
            except Exception as e:
                logging.warning(f"Excel weekly {user_id}: {e}")


async def check_and_send_daily():
    now_utc = datetime.now(timezone.utc)
    for user_id, time_str, tz_offset, send_excel in get_all_daily_users():
        try:
            hh, mm = map(int, time_str.split(":"))
        except Exception:
            continue
        user_now = now_utc + timedelta(hours=tz_offset)
        if user_now.hour != hh or user_now.minute != mm:
            continue
        target = (user_now - timedelta(days=1)).date()
        prev = target - timedelta(days=1)
        date_str = target.strftime("%Y-%m-%d")
        prev_str = prev.strftime("%Y-%m-%d")
        ops = get_operations_for_range(user_id, date_str, date_str)
        prev_ops = get_operations_for_range(user_id, prev_str, prev_str)
        prev_exp, prev_inc = 0.0, 0.0
        for op_type, amount, *_ in prev_ops:
            if op_type == "expense":
                prev_exp += amount
            else:
                prev_inc += amount
        if not ops:
            text = (f"📅 <b>Отчёт за {target.strftime('%d.%m.%Y')}</b>\n\n"
                    f"За этот день не было ни трат, ни зачислений.")
        else:
            agg = _aggregate(ops)
            text = _make_daily_report(
                date_str, agg,
                prev_totals=(prev_exp, prev_inc),
                expense_order=get_categories(user_id, "expense"),
                income_order=get_categories(user_id, "income"))
        try:
            await bot.send_message(user_id, text, parse_mode="HTML")
        except Exception as e:
            logging.warning(f"Не отправить daily {user_id}: {e}")
        if send_excel:
            try:
                await _send_excel(user_id, date_str, date_str,
                                  title=f"📎 Excel за {target.strftime('%d.%m.%Y')}",
                                  reason="daily",
                                  filename=f"finance_day_{target.strftime('%Y%m%d')}.xlsx")
            except Exception as e:
                logging.warning(f"Excel daily {user_id}: {e}")


async def check_and_send_reminders():
    now_utc = datetime.now(timezone.utc)
    for row in get_all_reminder_users():
        (user_id, interval_h, last_sent, silent_from,
         silent_to, tz_offset) = row
        user_now = now_utc + timedelta(hours=tz_offset)
        if silent_from > silent_to:
            in_silent = user_now.hour >= silent_from or user_now.hour < silent_to
        else:
            in_silent = silent_from <= user_now.hour < silent_to
        if in_silent:
            continue
        last_dt = None
        if last_sent:
            try:
                last_dt = datetime.strptime(last_sent, "%Y-%m-%d %H:%M:%S")
            except Exception:
                last_dt = None
        last_op_str = get_last_operation_dt(user_id)
        last_op_dt = None
        if last_op_str:
            try:
                last_op_dt = datetime.strptime(last_op_str, "%Y-%m-%d %H:%M:%S")
            except Exception:
                last_op_dt = None
        candidates = [d for d in (last_dt, last_op_dt) if d is not None]
        last_event = max(candidates) if candidates else None
        now_naive = now_utc.replace(tzinfo=None)
        if last_event and (now_naive - last_event) < timedelta(hours=interval_h):
            continue
        try:
            await bot.send_message(
                user_id,
                "🔔 <b>Напоминание</b>\n\n"
                "Не забудь записать траты или зачисления за сегодня.",
                parse_mode="HTML")
            update_reminder_last_sent(user_id,
                                      now_naive.strftime("%Y-%m-%d %H:%M:%S"))
        except Exception as e:
            logging.warning(f"Напомнить не удалось {user_id}: {e}")


async def check_and_send_autoexport():
    now_utc = datetime.now(timezone.utc)
    for user_id, weekday, time_str, period, tz_offset in get_all_autoexport_users():
        try:
            hh, mm = map(int, time_str.split(":"))
        except Exception:
            continue
        user_now = now_utc + timedelta(hours=tz_offset)
        if user_now.weekday() != weekday:
            continue
        if user_now.hour != hh or user_now.minute != mm:
            continue
        if period == "week":
            start = (user_now - timedelta(days=user_now.weekday())).replace(
                hour=0, minute=0, second=0, microsecond=0)
            end = user_now
            title = (f"📤 Автоэкспорт за неделю "
                     f"{start.strftime('%d.%m')} — {end.strftime('%d.%m.%Y')}")
            filename = f"autoexport_week_{start.strftime('%Y%m%d')}.xlsx"
        elif period == "month":
            start = user_now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
            end = user_now
            title = f"📤 Автоэкспорт за {start.strftime('%m.%Y')}"
            filename = f"autoexport_{start.strftime('%Y-%m')}.xlsx"
        else:
            earliest = get_earliest_operation_date(user_id)
            start = earliest or user_now
            if isinstance(start, str):
                start = datetime.strptime(start, "%Y-%m-%d")
            end = user_now
            title = "📤 Автоэкспорт за всё время"
            filename = f"autoexport_all_{user_now.strftime('%Y%m%d')}.xlsx"
        try:
            await _send_excel(user_id, start.strftime("%Y-%m-%d"),
                              end.strftime("%Y-%m-%d"),
                              title=title, reason="autoexport", filename=filename)
        except Exception as e:
            logging.warning(f"Автоэкспорт {user_id}: {e}")


# ============================================================
#                        ЗАПУСК
# ============================================================
async def main():
    init_db()
    init_settings_table()
    init_debts_tables()
    init_invest_tables()
    init_profit_table()
    init_export_history_table()
    init_accounts_table()
    init_view_log_table()
    init_access_requests_table()
    init_categories_table()
    migrate_existing_users()

    scheduler.add_job(check_and_send_weekly, "cron", minute="*")
    scheduler.add_job(check_and_send_daily, "cron", minute="*")
    scheduler.add_job(check_and_send_reminders, "cron", minute="*")
    scheduler.add_job(check_and_send_autoexport, "cron", minute="*")
    scheduler.start()

    print("Бот запущен... Планировщик работает.")
    try:
        await dp.start_polling(bot)
    finally:
        scheduler.shutdown(wait=False)


if __name__ == "__main__":
    asyncio.run(main())
